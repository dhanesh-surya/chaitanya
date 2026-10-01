from django.test import Client, TestCase
from django.urls import reverse

from accounts.models import AdminUser, Student


class ManageStudentsProgramFilterTestCase(TestCase):
    def setUp(self):
        self.admin = AdminUser.objects.create(
            username='admin',
            password='adminpassword',
        )
        self.student_bsc = Student.objects.create(
            registration_no='BSC001',
            full_name='Amit Verma',
            email='amit@example.com',
            program_type='B.Sc. First Semester',
            course_name='B.Sc.',
        )
        self.student_ba = Student.objects.create(
            registration_no='BA001',
            full_name='Priya Singh',
            email='priya@example.com',
            program_type='B.A. First Semester',
            course_name='B.A.',
        )

    def _login_admin(self):
        session = self.client.session
        session['admin_user'] = 'admin'
        session.save()

    def test_default_students_page_is_blank_and_prompts_to_select_program(self):
        self._login_admin()
        response = self.client.get(reverse('manage_students'))
        self.assertEqual(response.status_code, 200)

        # By default, no students are loaded
        self.assertEqual(len(response.context['students']), 0)
        self.assertEqual(response.context['total_count'], 0)
        self.assertEqual(response.context['program_filter'], '')

        content = response.content.decode('utf-8')
        # "All Programs" must be removed
        self.assertNotIn('All Programs', content)
        # "-- Select Program --" option should be present
        self.assertIn('-- Select Program --', content)
        # Empty state message indicating no program selected
        self.assertIn('No Program Selected', content)
        self.assertIn('Please select a program from the dropdown above to see details.', content)

    def test_students_details_load_when_program_selected(self):
        self._login_admin()
        response = self.client.get(reverse('manage_students') + '?program=B.Sc.+First+Semester')
        self.assertEqual(response.status_code, 200)

        # Only B.Sc. students loaded
        self.assertEqual(len(response.context['students']), 1)
        self.assertEqual(response.context['students'][0].registration_no, 'BSC001')
        self.assertEqual(response.context['total_count'], 1)

        content = response.content.decode('utf-8')
        self.assertIn('Amit Verma', content)
        self.assertNotIn('Priya Singh', content)

    def test_export_csv_without_program_warns_and_redirects(self):
        self._login_admin()
        response = self.client.get(reverse('export_students_csv'))
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse('manage_students'))

    def test_export_csv_with_program(self):
        self._login_admin()
        response = self.client.get(reverse('export_students_csv') + '?program=B.Sc.+First+Semester')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv')
        content = response.content.decode('utf-8')
        self.assertIn('BSC001', content)
        self.assertIn('Amit Verma', content)
        self.assertNotIn('BA001', content)

    def test_admin_can_print_slip_for_any_student_enrollment(self):
        from admissions.models import StudentEnrollment

        enrollment = StudentEnrollment.objects.create(
            enrollment_no='CCP990001',
            reg_no='BSC001',
            student=self.student_bsc,
            full_name='Amit Verma',
            program_type='B.Sc. First Semester',
            is_submitted=True,
            status='Submitted',
        )

        # Login as admin, and also simulate another student's reg_no in session
        session = self.client.session
        session['admin_user'] = 'admin'
        session['reg_no'] = 'OTHER_STUDENT_999'
        session.save()

        # Print slip should successfully open and return 200
        response = self.client.get(reverse('enrollment_print', kwargs={'enrollment_no': enrollment.enrollment_no}))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8')
        self.assertIn('Amit Verma', content)
        self.assertIn('CCP990001', content)

    def test_manage_enrollments_view_and_filters(self):
        from admissions.models import StudentEnrollment

        e1 = StudentEnrollment.objects.create(
            enrollment_no='CCP990002',
            reg_no='BSC001',
            student=self.student_bsc,
            full_name='Amit Verma',
            program_type='B.Sc. First Semester',
            is_submitted=True,
            status='Submitted',
        )
        e2 = StudentEnrollment.objects.create(
            enrollment_no='CCP990003',
            reg_no='BA001',
            student=self.student_ba,
            full_name='Priya Singh',
            program_type='B.A. First Semester',
            is_submitted=True,
            status='Approved',
        )

        self._login_admin()
        # View all
        resp = self.client.get(reverse('manage_enrollments'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.context['enrollments']), 2)
        content = resp.content.decode('utf-8')
        self.assertIn('Manage Student Enrollments', content)
        self.assertIn('Amit Verma', content)
        self.assertIn('Priya Singh', content)

        # Filter by program
        resp_prog = self.client.get(reverse('manage_enrollments') + '?program=B.Sc.+First+Semester')
        self.assertEqual(resp_prog.status_code, 200)
        self.assertEqual(len(resp_prog.context['enrollments']), 1)
        self.assertEqual(resp_prog.context['enrollments'][0].reg_no, 'BSC001')

        # Filter by status
        resp_status = self.client.get(reverse('manage_enrollments') + '?status=Approved')
        self.assertEqual(resp_status.status_code, 200)
        self.assertEqual(len(resp_status.context['enrollments']), 1)
        self.assertEqual(resp_status.context['enrollments'][0].reg_no, 'BA001')

    def test_admin_single_enrollment_approve_and_reset(self):
        from admissions.models import StudentEnrollment

        e = StudentEnrollment.objects.create(
            reg_no='BSC001',
            student=self.student_bsc,
            full_name='Amit Verma',
            program_type='B.Sc. First Semester',
            is_submitted=True,
            status='Submitted',
        )
        self.assertFalse(self.student_bsc.is_verified)
        self.assertIsNone(e.enrollment_no)

        self._login_admin()
        # Approve enrollment
        resp_approve = self.client.post(
            reverse('update_enrollment_status', kwargs={'pk': e.pk}),
            {'status': 'Approved'}
        )
        self.assertEqual(resp_approve.status_code, 302)

        e.refresh_from_db()
        self.assertEqual(e.status, 'Approved')
        self.assertTrue(e.is_submitted)
        self.assertIsNotNone(e.enrollment_no)  # Auto-generated on approval!
        self.student_bsc.refresh_from_db()
        self.assertTrue(self.student_bsc.is_verified)

        # Reset to Draft
        resp_draft = self.client.post(
            reverse('update_enrollment_status', kwargs={'pk': e.pk}),
            {'status': 'Draft'}
        )
        self.assertEqual(resp_draft.status_code, 302)
        e.refresh_from_db()
        self.assertEqual(e.status, 'Draft')
        self.assertFalse(e.is_submitted)

    def test_admin_bulk_approve_enrollments(self):
        from admissions.models import StudentEnrollment

        e1 = StudentEnrollment.objects.create(
            reg_no='BSC001',
            student=self.student_bsc,
            full_name='Amit Verma',
            program_type='B.Sc. First Semester',
            is_submitted=True,
            status='Submitted',
        )
        e2 = StudentEnrollment.objects.create(
            reg_no='BA001',
            student=self.student_ba,
            full_name='Priya Singh',
            program_type='B.A. First Semester',
            is_submitted=True,
            status='Submitted',
        )

        self._login_admin()
        # Bulk approve
        post_data = {
            'enrollment_ids': [e1.pk, e2.pk],
            'status': 'Approved',
        }
        resp = self.client.post(reverse('bulk_update_enrollment_status'), post_data)
        self.assertEqual(resp.status_code, 302)

        e1.refresh_from_db()
        e2.refresh_from_db()
        self.assertEqual(e1.status, 'Approved')
        self.assertEqual(e2.status, 'Approved')
        self.assertTrue(e1.is_submitted)
        self.assertTrue(e2.is_submitted)
        self.assertIsNotNone(e1.enrollment_no)
        self.assertIsNotNone(e2.enrollment_no)

        self.student_bsc.refresh_from_db()
        self.student_ba.refresh_from_db()
        self.assertTrue(self.student_bsc.is_verified)
        self.assertTrue(self.student_ba.is_verified)


class ExportEnrollmentsExcelTestCase(TestCase):
    def setUp(self):
        import json
        from datetime import date
        from accounts.models import AdminUser, Student
        from admissions.models import StudentAdmission, StudentEnrollment

        self.client = Client()
        self.admin = AdminUser.objects.create(
            username='export_admin',
            password='admin_secret_pass',
        )

        self.student_m = Student.objects.create(
            registration_no='REG_M_01',
            full_name='Rohan Sahu',
            mobile='9876543210',
            password='pass1',
        )
        self.adm_m = StudentAdmission.objects.create(
            reg_no='REG_M_01',
            application_no='APP2026M01',
            full_name='Rohan Sahu',
            father_name='Ramesh Sahu',
            mother_name='Savitri Sahu',
            gender='Male',
            dob=date(2003, 4, 15),
            category='OBC',
            medium='Hindi',
            perm_village='Pamgarh',
            perm_city='Pamgarh',
            perm_district='Janjgir-Champa',
            perm_state='Chhattisgarh',
            perm_pin_code='495554',
            subject='Hindi Language, History',
        )
        self.enr_m = StudentEnrollment.objects.create(
            enrollment_no='CCP26059001',
            reg_no='REG_M_01',
            student=self.student_m,
            admission=self.adm_m,
            full_name='Rohan Sahu',
            father_name='Ramesh Sahu',
            mother_name='Savitri Sahu',
            gender='Male',
            dob=date(2003, 4, 15),
            category='OBC',
            medium='Hindi',
            mobile='9876543210',
            perm_village='Pamgarh',
            perm_city='Pamgarh',
            perm_district='Janjgir-Champa',
            perm_state='Chhattisgarh',
            perm_pin_code='495554',
            program_type='B.A. First Semester',
            status='Approved',
            is_submitted=True,
            selected_courses_json=json.dumps([
                {'code': 'HNSC-01', 'name': 'Hindi — Hindi Sahitya Ka Itihas'},
                {'code': 'HISC-01', 'name': 'History — Ancient Indian History'},
            ]),
        )

        self.student_f = Student.objects.create(
            registration_no='REG_F_01',
            full_name='Kavita Patel',
            mobile='9876543211',
            password='pass2',
        )
        self.enr_f = StudentEnrollment.objects.create(
            enrollment_no='CCP26059002',
            reg_no='REG_F_01',
            student=self.student_f,
            full_name='Kavita Patel',
            father_name='Dinesh Patel',
            mother_name='Geeta Patel',
            gender='Female',
            dob=date(2004, 10, 22),
            category='GEN',
            medium='English',
            mobile='9876543211',
            perm_village='Bilaspur',
            perm_city='Bilaspur',
            perm_district='Bilaspur',
            perm_state='Chhattisgarh',
            perm_pin_code='495001',
            program_type='B.Sc. First Semester',
            status='Submitted',
            is_submitted=True,
            selected_courses_json=json.dumps({
                'courses': [
                    {'code': 'CHSC-01T', 'name': 'Chemistry — Fundamental Chemistry-I'},
                    {'code': 'BOSC-01T', 'name': 'Botany — Elementary Botany'},
                ]
            }),
        )

    def _login_admin(self):
        session = self.client.session
        session['admin_user'] = self.admin.username
        session['is_admin_logged_in'] = True
        session.save()

    def test_export_requires_admin_login(self):
        resp = self.client.get(reverse('export_enrollments_excel'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn('/admin/login/', resp.url)

    def test_export_enrollments_excel_headers_and_data(self):
        from io import BytesIO
        import openpyxl

        self._login_admin()
        resp = self.client.get(reverse('export_enrollments_excel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment; filename="enrollments_', resp['Content-Disposition'])

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.title, 'Enrollments')

        # Check headers (row 1)
        expected_headers = [
            'AddmissionNo',
            'Univ_EnrolNo',
            'Stud_nm',
            'FatherName',
            'MotherName',
            'Medium',
            'Category',
            'Gender (MALE-1 ,FEMALE-0)',
            'DOB (MM/DD/YYYY)',
            'Address',
            'Mobile',
            'CLASS NAME',
            'SUBJECT CODE',
            'SUBJECTS',
        ]
        actual_headers = [ws.cell(row=1, column=col).value for col in range(1, len(expected_headers) + 1)]
        self.assertEqual(actual_headers, expected_headers)

        # Check data rows (rows 2 and 3)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertGreaterEqual(len(rows), 2)

        # Find row for Rohan Sahu (Male)
        rohan_row = next(r for r in rows if r[2] == 'Rohan Sahu')
        self.assertEqual(rohan_row[0], 'APP2026M01')  # AddmissionNo
        self.assertEqual(rohan_row[1], 'CCP26059001')  # Univ_EnrolNo
        self.assertEqual(rohan_row[3], 'Ramesh Sahu')  # FatherName
        self.assertEqual(rohan_row[4], 'Savitri Sahu')  # MotherName
        self.assertEqual(rohan_row[5], 'Hindi')  # Medium
        self.assertEqual(rohan_row[6], 'OBC')  # Category
        self.assertEqual(rohan_row[7], 1)  # Gender: MALE-1
        self.assertEqual(rohan_row[8], '04/15/2003')  # DOB: MM/DD/YYYY
        self.assertIn('Pamgarh', rohan_row[9])  # Address
        self.assertEqual(rohan_row[10], '9876543210')  # Mobile
        self.assertEqual(rohan_row[11], 'B.A. First Semester')  # CLASS NAME
        self.assertIn('HNSC-01', rohan_row[12])  # SUBJECT CODE
        self.assertIn('HISC-01', rohan_row[12])
        self.assertIn('Hindi Sahitya Ka Itihas', rohan_row[13])  # SUBJECTS
        self.assertIn('Ancient Indian History', rohan_row[13])

        # Find row for Kavita Patel (Female)
        kavita_row = next(r for r in rows if r[2] == 'Kavita Patel')
        self.assertEqual(kavita_row[1], 'CCP26059002')  # Univ_EnrolNo
        self.assertEqual(kavita_row[7], 0)  # Gender: FEMALE-0
        self.assertEqual(kavita_row[8], '10/22/2004')  # DOB: MM/DD/YYYY
        self.assertEqual(kavita_row[11], 'B.Sc. First Semester')  # CLASS NAME
        self.assertIn('CHSC-01T', kavita_row[12])  # SUBJECT CODE
        self.assertIn('BOSC-01T', kavita_row[12])
        self.assertIn('Fundamental Chemistry-I', kavita_row[13])  # SUBJECTS

    def test_export_enrollments_excel_filtering(self):
        from io import BytesIO
        import openpyxl

        self._login_admin()
        # Filter by program
        resp = self.client.get(reverse('export_enrollments_excel') + '?program=B.A.+First+Semester')
        self.assertEqual(resp.status_code, 200)

        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # Should only contain B.A. First Semester
        self.assertTrue(all(r[11] == 'B.A. First Semester' for r in rows))
        self.assertTrue(any(r[2] == 'Rohan Sahu' for r in rows))
        self.assertFalse(any(r[2] == 'Kavita Patel' for r in rows))



