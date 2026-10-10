import csv
import json
import re

from django.contrib import messages
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from accounts.models import AdminUser, Student
from accounts.utils import (
    admin_login_required,
    generate_secure_password,
    is_valid_aadhaar,
    is_valid_email,
    is_valid_mobile,
)
from admissions.models import StudentAdmission, StudentEnrollment
from courses.models import ProgramCourse
from admissions.services import (
    get_print_context,
    parse_selected_subjects,
    parse_selected_subjects_payload,
)
from courses.subject_groups import (
    format_bsc_group_display_name,
    get_bsc_subject_group_sections,
    is_bsc_program,
)
from courses.utils import get_program_names

from .merit_list import (
    MERIT_EXPORT_HEADERS,
    STATUS_FILTER_CHOICES,
    build_merit_list_groups,
    build_merit_list_workbook,
    get_merit_program_choices,
    iter_merit_export_rows,
    merit_export_filename,
)


def _students_filter_params(request):
    """Read list filters from POST (actions) or GET (page load / filter form)."""
    if request.method == 'POST':
        source = request.POST
    else:
        source = request.GET

    program = (source.get('program') or '').strip()
    if program == 'ALL':
        program = ''
    verified = (source.get('verified') or 'ALL').strip() or 'ALL'
    group = (source.get('group') or 'ALL').strip() or 'ALL'
    return {
        'search': source.get('search', '').strip(),
        'program': program,
        'verified': verified,
        'group': group,
    }


def _subject_group_filter_choices(program_type):
    """B.Sc. subject groups for the selected program (empty when not B.Sc.)."""
    if not is_bsc_program(program_type):
        return []
    choices = []
    for section in get_bsc_subject_group_sections(program_type):
        for group in section['groups']:
            choices.append({
                'key': group['key'],
                'label': group.get('full_name') or f"{section['heading']} — {group['label']}",
            })
    return choices


def _normalize_group_filter(group_filter, program_type):
    """Drop invalid group keys when program is not B.Sc. or group is unknown."""
    group_filter = (group_filter or 'ALL').strip() or 'ALL'
    if group_filter == 'ALL':
        return 'ALL'
    choices = _subject_group_filter_choices(program_type)
    if not choices:
        return 'ALL'
    valid = {item['key'] for item in choices}
    if group_filter == 'NONE' or group_filter in valid:
        return group_filter
    return 'ALL'


def _admission_bsc_group_key(admission):
    if not admission:
        return ''
    _, group_key = parse_selected_subjects_payload(admission.selected_subjects_json)
    return (group_key or '').strip()


def _filter_students_by_group(students, group_filter='ALL'):
    if not group_filter or group_filter == 'ALL':
        return students
    result = []
    for student in students:
        key = _admission_bsc_group_key(getattr(student, 'latest_admission', None))
        if group_filter == 'NONE':
            if not key:
                result.append(student)
        elif key == group_filter:
            result.append(student)
    return result


def _students_url(params=None, edit_pk=None):
    from urllib.parse import urlencode
    from django.urls import reverse

    query = dict(params or {})
    if edit_pk:
        query['edit'] = edit_pk
    base = reverse('manage_students')
    if not query:
        return base
    return f'{base}?{urlencode(query)}'


def _get_students_queryset(search='', program_filter='', verified_filter='ALL'):
    if not program_filter and not search:
        return Student.objects.none()
    students = Student.objects.all()
    if search:
        students = students.filter(
            Q(full_name__icontains=search)
            | Q(email__icontains=search)
            | Q(mobile__icontains=search)
            | Q(registration_no__icontains=search)
            | Q(aadhaar__icontains=search)
        )
    if program_filter:
        students = students.filter(program_type=program_filter)
    if verified_filter == 'YES':
        students = students.filter(is_verified=True)
    elif verified_filter == 'NO':
        students = students.filter(is_verified=False)
    return students.order_by('-created_date')


def _export_course_name_only(label):
    """Strip B.Sc. department prefix; keep paper/course name only."""
    name = (label or '').strip()
    if not name:
        return ''
    if ' — ' in name:
        _department, course_name = name.split(' — ', 1)
        if course_name.strip():
            return course_name.strip()
    return name


def _format_selected_courses_for_export(student):
    """Paper names from the student's latest admission, for CSV export."""
    admission = getattr(student, 'latest_admission', None)
    if admission:
        subjects = parse_selected_subjects(admission)
        if subjects:
            names = [
                _export_course_name_only(s.get('name'))
                for s in subjects
                if isinstance(s, dict) and _export_course_name_only(s.get('name'))
            ]
            if names:
                return '; '.join(names)
        if (admission.subject or '').strip():
            names = [
                _export_course_name_only(part)
                for part in admission.subject.split(',')
                if _export_course_name_only(part)
            ]
            if names:
                return '; '.join(names)
    return _export_course_name_only(student.course_name)


def _admission_dsc_course_labels(admission):
    """DSC theory course labels from selected subjects (department — paper name)."""
    if not admission:
        return []
    subjects = parse_selected_subjects(admission)
    labels = []
    seen = set()
    for item in subjects:
        if not isinstance(item, dict):
            continue
        if (item.get('type2') or '').strip().upper() != 'DSC':
            continue
        type1 = (item.get('type1') or '').strip().lower()
        if 'practical' in type1 or 'lab' in type1:
            continue
        full = (item.get('name') or '').strip()
        if not full:
            continue
        key = full.lower()
        if key in seen:
            continue
        seen.add(key)
        labels.append(full)
    return labels


def _attach_admissions(students):
    reg_nos = [s.registration_no for s in students]
    admission_map = {}
    enrollment_map = {}
    nep_map = {}
    if reg_nos:
        for adm in (
            StudentAdmission.objects.filter(reg_no__in=reg_nos)
            .order_by('-submitted_date', '-created_date')
        ):
            if adm.reg_no not in admission_map:
                admission_map[adm.reg_no] = adm
        for enr in (
            StudentEnrollment.objects.filter(reg_no__in=reg_nos)
            .order_by('-submitted_date', '-created_at')
        ):
            if enr.reg_no not in enrollment_map:
                enrollment_map[enr.reg_no] = enr
        from admissions.models import NepUgAdmissionEnrollment
        for nep in (
            NepUgAdmissionEnrollment.objects.filter(reg_no__in=reg_nos)
            .order_by('-submitted_date', '-created_at')
        ):
            if nep.reg_no not in nep_map:
                nep_map[nep.reg_no] = nep

    for student in students:
        admission = admission_map.get(student.registration_no)
        enrollment = enrollment_map.get(student.registration_no)
        nep = nep_map.get(student.registration_no)
        student.latest_admission = admission
        student.latest_enrollment = enrollment
        student.latest_nep_ug = nep
        student.father_name = (
            (admission.father_name if admission and admission.father_name else '')
            or (enrollment.father_name if enrollment and enrollment.father_name else '')
            or (nep.father_name if nep and nep.father_name else '')
        )
        student.mother_name = (
            (admission.mother_name if admission and admission.mother_name else '')
            or (enrollment.mother_name if enrollment and enrollment.mother_name else '')
            or (nep.mother_name if nep and nep.mother_name else '')
        )
        group_key = _admission_bsc_group_key(admission)
        student.bsc_group_key = group_key
        student.bsc_group_label = (
            format_bsc_group_display_name(group_key, student.program_type)
            if group_key
            else ''
        )
        student.dsc_course_names = _admission_dsc_course_labels(admission)
        student.dsc_courses_display = '; '.join(student.dsc_course_names)
    return students


def _program_group_label(program_type):
    label = (program_type or '').strip()
    return label or 'Not Assigned'


def _build_verified_students_by_program():
    """Group verified students by program (class) for admin dashboard."""
    verified = list(
        Student.objects.filter(is_verified=True).order_by('program_type', 'full_name', 'registration_no')
    )
    _attach_admissions(verified)

    groups = {}
    for student in verified:
        program = _program_group_label(student.program_type)
        groups.setdefault(program, []).append(student)

    return [
        {
            'program': program,
            'count': len(students),
            'students': students,
        }
        for program, students in sorted(
            groups.items(),
            key=lambda item: (item[0] == 'Not Assigned', item[0].lower()),
        )
    ]


_REJECTABLE_ADMISSION_STATUSES = ('Submitted', 'Pending', 'Approved')


def _get_latest_admission_for_student(student):
    return (
        StudentAdmission.objects.filter(reg_no=student.registration_no)
        .order_by('-submitted_date', '-created_date')
        .first()
    )


def _reject_student_application(student):
    admission = _get_latest_admission_for_student(student)
    if not admission or admission.status not in _REJECTABLE_ADMISSION_STATUSES:
        return False
    admission.status = 'Rejected'
    admission.save(update_fields=['status'])
    return True


def _validate_student_form(post, student=None):
    errors = []
    full_name = post.get('full_name', '').strip()
    email = post.get('email', '').strip().lower()
    mobile = re.sub(r'\D', '', post.get('mobile', '').strip())
    aadhaar = re.sub(r'\D', '', post.get('aadhaar', '').strip())
    program_type = post.get('program_type', '').strip()

    if not full_name:
        errors.append('Full name is required.')
    if email and not is_valid_email(email):
        errors.append('Invalid email format.')
    if mobile and not is_valid_mobile(mobile):
        errors.append('Invalid mobile number.')
    if aadhaar and not is_valid_aadhaar(aadhaar):
        errors.append('Invalid Aadhaar number.')

    if email:
        qs = Student.objects.filter(email__iexact=email)
        if student:
            qs = qs.exclude(pk=student.pk)
        if qs.exists():
            errors.append('Another student already uses this email.')

    if mobile:
        qs = Student.objects.filter(mobile=mobile)
        if student:
            qs = qs.exclude(pk=student.pk)
        if qs.exists():
            errors.append('Another student already uses this mobile.')

    if aadhaar:
        qs = Student.objects.filter(aadhaar=aadhaar)
        if student:
            qs = qs.exclude(pk=student.pk)
        if qs.exists():
            errors.append('Another student already uses this Aadhaar.')

    return errors, {
        'full_name': full_name,
        'email': email,
        'mobile': mobile,
        'aadhaar': aadhaar,
        'program_type': program_type,
        'course_name': post.get('course_name', '').strip(),
        'is_verified': post.get('is_verified') == 'on',
    }


@require_http_methods(['GET', 'POST'])
def admin_login(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '').strip()
        if AdminUser.objects.filter(username=username, password=password).exists():
            request.session['admin_user'] = username
            return redirect('admin_dashboard')
        messages.error(request, 'Invalid username or password.')
    return render(request, 'admin_panel/login.html')


def admin_logout(request):
    request.session.pop('admin_user', None)
    request.session.pop('courses_manage_unlocked', None)
    return redirect('admin_login')


@admin_login_required
def admin_dashboard(request):
    from admissions.models import NepUgAdmissionEnrollment, StudentEnrollment

    show_verified = request.GET.get('show_verified') == '1'
    verified_students = Student.objects.filter(is_verified=True).count()
    stats = {
        'total_admissions': StudentAdmission.objects.count(),
        'pending': StudentAdmission.objects.filter(status='Pending').count(),
        'approved': StudentAdmission.objects.filter(status='Approved').count(),
        'rejected': StudentAdmission.objects.filter(status='Rejected').count(),
        'submitted': StudentAdmission.objects.filter(status='Submitted').count(),
        'students': Student.objects.count(),
        'verified_students': verified_students,
        'total_enrollments': StudentEnrollment.objects.count(),
        'submitted_enrollments': StudentEnrollment.objects.filter(is_submitted=True).count(),
        'approved_enrollments': StudentEnrollment.objects.filter(status='Approved').count(),
        'total_nepug': NepUgAdmissionEnrollment.objects.count(),
        'submitted_nepug': NepUgAdmissionEnrollment.objects.filter(is_submitted=True).count(),
        'approved_nepug': NepUgAdmissionEnrollment.objects.filter(status='Approved').count(),
    }
    recent = StudentAdmission.objects.order_by('-submitted_date', '-created_date')[:20]
    recent_enrollments = StudentEnrollment.objects.order_by('-submitted_date', '-created_at')[:20]
    return render(request, 'admin_panel/dashboard.html', {
        'stats': stats,
        'recent': recent,
        'recent_enrollments': recent_enrollments,
        'show_verified': show_verified,
        'verified_by_program': _build_verified_students_by_program() if show_verified else [],
    })


@admin_login_required
def manage_students(request):
    params = _students_filter_params(request)
    search = params['search']
    program_filter = params['program']
    verified_filter = params['verified']
    group_filter = _normalize_group_filter(params['group'], program_filter)
    params['group'] = group_filter
    edit_pk = request.GET.get('edit', '').strip()

    students = _filter_students_by_group(
        _attach_admissions(
            list(_get_students_queryset(search, program_filter, verified_filter))
        ),
        group_filter,
    )

    edit_student = None
    if edit_pk:
        edit_student = Student.objects.filter(pk=edit_pk).first()
        if edit_student:
            _attach_admissions([edit_student])

    program_types = get_program_names(active_only=False)
    subject_group_choices = _subject_group_filter_choices(program_filter)
    show_group_filter = bool(subject_group_choices)
    reset_password_display = request.session.pop('reset_password_display', None)
    bulk_reset_passwords_display = request.session.pop('bulk_reset_passwords_display', None)

    return render(request, 'admin_panel/students.html', {
        'students': students,
        'search': search,
        'program_filter': program_filter,
        'verified_filter': verified_filter,
        'group_filter': group_filter,
        'program_types': program_types,
        'subject_group_choices': subject_group_choices,
        'show_group_filter': show_group_filter,
        'edit_student': edit_student,
        'filter_params': params,
        'total_count': len(students),
        'reset_password_display': reset_password_display,
        'bulk_reset_passwords_display': bulk_reset_passwords_display,
    })


@admin_login_required
@require_http_methods(['POST'])
def edit_student(request, pk):
    student = get_object_or_404(Student, pk=pk)
    errors, data = _validate_student_form(request.POST, student=student)
    params = _students_filter_params(request)

    if errors:
        for err in errors:
            messages.error(request, err)
        return redirect(_students_url(params, edit_pk=pk))

    student.full_name = data['full_name']
    student.email = data['email'] or None
    student.mobile = data['mobile'] or None
    student.aadhaar = data['aadhaar'] or None
    student.program_type = data['program_type']
    student.course_name = data['course_name']
    student.is_verified = data['is_verified']
    student.save()
    messages.success(request, f'Student {student.registration_no} updated successfully.')
    return redirect(_students_url(params))


@admin_login_required
@require_http_methods(['POST'])
def delete_student(request, pk):
    student = get_object_or_404(Student, pk=pk)
    params = _students_filter_params(request)

    if StudentAdmission.objects.filter(reg_no=student.registration_no, status='Submitted').exists():
        messages.error(
            request,
            f'Cannot delete {student.registration_no} — reject the submitted application first.',
        )
        return redirect(_students_url(params))

    reg_no = student.registration_no
    StudentAdmission.objects.filter(reg_no=reg_no).delete()
    student.delete()
    messages.success(request, f'Student {reg_no} deleted.')
    return redirect(_students_url(params))


@admin_login_required
@require_http_methods(['POST'])
def reset_student_password(request, pk):
    student = get_object_or_404(Student, pk=pk)
    params = _students_filter_params(request)
    new_password = generate_secure_password()
    student.password = new_password
    student.save(update_fields=['password'])
    request.session['reset_password_display'] = {
        'registration_no': student.registration_no,
        'password': new_password,
    }
    messages.success(request, f'Password reset for {student.registration_no}.')
    return redirect(_students_url(params))


def _unverify_and_reassign_student(student):
    """
    Mark student as unverified and reset their latest admission status to 'Draft'
    so they can completely edit their form.
    """
    student.is_verified = False
    student.save(update_fields=['is_verified'])

    from admissions.models import StudentAdmission
    admission = (
        StudentAdmission.objects.filter(reg_no=student.registration_no)
        .order_by('-submitted_date', '-created_date')
        .first()
    )
    if admission:
        admission.status = 'Draft'
        admission.is_submitted = False
        admission.is_approved = False
        admission.save(update_fields=['status', 'is_submitted', 'is_approved'])


def _verify_student(student):
    """
    Mark student as verified and approve their latest admission.
    """
    student.is_verified = True
    student.save(update_fields=['is_verified'])

    from admissions.models import StudentAdmission
    admission = (
        StudentAdmission.objects.filter(reg_no=student.registration_no)
        .order_by('-submitted_date', '-created_date')
        .first()
    )
    if admission:
        admission.status = 'Approved'
        admission.is_submitted = True
        admission.is_approved = True
        admission.save(update_fields=['status', 'is_submitted', 'is_approved'])


@admin_login_required
@require_http_methods(['POST'])
def toggle_student_verified(request, pk):
    student = get_object_or_404(Student, pk=pk)
    params = _students_filter_params(request)
    if student.is_verified:
        _unverify_and_reassign_student(student)
        messages.success(request, f'{student.registration_no} marked as unverified and reassigned to Draft status.')
    else:
        _verify_student(student)
        messages.success(request, f'{student.registration_no} marked as verified and application approved.')
    return redirect(_students_url(params))


@admin_login_required
@require_http_methods(['POST'])
def bulk_student_action(request):
    valid_actions = ('verify', 'unverify', 'reset_password', 'reject_application', 'delete')
    params = _students_filter_params(request)
    raw_ids = request.POST.getlist('student_ids')
    student_ids = []
    for raw_id in raw_ids:
        try:
            student_ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue

    action = request.POST.get('action', '').strip()
    if not student_ids:
        messages.error(request, 'Select at least one student.')
        return redirect(_students_url(params))
    if action not in valid_actions:
        messages.error(request, 'Choose a valid action.')
        return redirect(_students_url(params))

    students = list(Student.objects.filter(pk__in=student_ids))
    if not students:
        messages.error(request, 'No matching students found.')
        return redirect(_students_url(params))

    if action == 'verify':
        for student in students:
            _verify_student(student)
        messages.success(
            request,
            f'Marked {len(students)} student{"s" if len(students) != 1 else ""} as verified and approved their applications.',
        )
    elif action == 'unverify':
        for student in students:
            _unverify_and_reassign_student(student)
        messages.success(
            request,
            f'Marked {len(students)} student{"s" if len(students) != 1 else ""} as not verified and reassigned to Draft status.',
        )
    elif action == 'reset_password':
        reset_list = []
        for student in students:
            new_password = generate_secure_password()
            student.password = new_password
            student.save(update_fields=['password'])
            reset_list.append({
                'registration_no': student.registration_no,
                'password': new_password,
            })
        request.session['bulk_reset_passwords_display'] = reset_list
        messages.success(
            request,
            f'Reset password for {len(reset_list)} student{"s" if len(reset_list) != 1 else ""}.',
        )
    elif action == 'reject_application':
        rejected = 0
        skipped = []
        for student in students:
            if _reject_student_application(student):
                rejected += 1
            else:
                skipped.append(student.registration_no)
        if rejected:
            messages.success(
                request,
                f'Rejected {rejected} application{"s" if rejected != 1 else ""}. '
                'Those students can now be deleted.',
            )
        if skipped:
            messages.warning(
                request,
                'Skipped '
                f'{len(skipped)} student{"s" if len(skipped) != 1 else ""} with no rejectable application: '
                + ', '.join(skipped),
            )
        if not rejected and not skipped:
            messages.error(request, 'No applications were rejected.')
    elif action == 'delete':
        deleted = 0
        skipped = []
        for student in students:
            if StudentAdmission.objects.filter(
                reg_no=student.registration_no,
                status='Submitted',
            ).exists():
                skipped.append(student.registration_no)
                continue
            reg_no = student.registration_no
            StudentAdmission.objects.filter(reg_no=reg_no).delete()
            student.delete()
            deleted += 1
        if deleted:
            messages.success(
                request,
                f'Deleted {deleted} student{"s" if deleted != 1 else ""}.',
            )
        if skipped:
            messages.warning(
                request,
                'Skipped '
                f'{len(skipped)} student{"s" if len(skipped) != 1 else ""} with submitted applications '
                '(reject the application first): '
                + ', '.join(skipped),
            )
        if not deleted and not skipped:
            messages.error(request, 'No students were deleted.')

    return redirect(_students_url(params))


def _merit_list_params(request):
    program = request.GET.get('program', 'ALL').strip() or 'ALL'
    status = request.GET.get('status', 'applied').strip() or 'applied'
    valid_statuses = {key for key, _ in STATUS_FILTER_CHOICES}
    if status not in valid_statuses:
        status = 'applied'
    return {'program': program, 'status': status}


@admin_login_required
def merit_list(request):
    params = _merit_list_params(request)
    program_filter = params['program']
    status_filter = params['status']
    groups = build_merit_list_groups(program_filter, status_filter)
    total_applications = sum(group['count'] for group in groups)

    return render(request, 'admin_panel/merit_list.html', {
        'groups': groups,
        'program_filter': program_filter,
        'status_filter': status_filter,
        'program_choices': get_merit_program_choices(status_filter),
        'status_choices': STATUS_FILTER_CHOICES,
        'total_applications': total_applications,
    })


@admin_login_required
def export_merit_list_csv(request):
    params = _merit_list_params(request)
    groups = build_merit_list_groups(params['program'], params['status'])

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = (
        f'attachment; filename="{merit_export_filename(params["program"], "csv")}"'
    )
    writer = csv.writer(response)
    writer.writerow(MERIT_EXPORT_HEADERS)
    for row in iter_merit_export_rows(groups):
        writer.writerow(row)
    return response


@admin_login_required
def export_merit_list_excel(request):
    params = _merit_list_params(request)
    groups = build_merit_list_groups(params['program'], params['status'])
    workbook = build_merit_list_workbook(
        groups,
        program_filter=params['program'],
        status_filter=params['status'],
    )

    response = HttpResponse(
        workbook.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="{merit_export_filename(params["program"], "xlsx")}"'
    )
    return response


@admin_login_required
def export_students_csv(request):
    params = _students_filter_params(request)
    if not params['program'] and not params['search']:
        messages.warning(request, 'Please select a program before exporting CSV.')
        return redirect('manage_students')
    group_filter = _normalize_group_filter(params['group'], params['program'])
    students = _filter_students_by_group(
        _attach_admissions(
            list(_get_students_queryset(params['search'], params['program'], params['verified']))
        ),
        group_filter,
    )

    response = HttpResponse(content_type='text/csv')
    filename = f"Students_{params['program'] or 'All'}.csv".replace(' ', '_').replace('/', '_')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow([
        'Registration No', 'Full Name', "Father's Name", "Mother's Name", 'Email', 'Mobile', 'Aadhaar',
        'Program', 'Subject Group', 'DSC Courses', 'Course', 'Selected Courses / Paper Name',
        'Application No', 'Verified', 'Registered On',
    ])
    for s in students:
        admission = getattr(s, 'latest_admission', None)
        father = getattr(s, 'father_name', '') or (admission.father_name if admission else '')
        mother = getattr(s, 'mother_name', '') or (admission.mother_name if admission else '')
        writer.writerow([
            s.registration_no,
            s.full_name,
            father,
            mother,
            s.email or '',
            s.mobile or '',
            s.aadhaar or '',
            s.program_type,
            getattr(s, 'bsc_group_label', '') or '',
            getattr(s, 'dsc_courses_display', '') or '',
            s.course_name,
            _format_selected_courses_for_export(s),
            (admission.application_no if admission else '') or '',
            'Yes' if s.is_verified else 'No',
            s.created_date.strftime('%d-%m-%Y %H:%M') if s.created_date else '',
        ])
    return response


@admin_login_required
def export_students_excel(request):
    """Export student list as an Excel (.xlsx) file with full details including Father & Mother name."""
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    params = _students_filter_params(request)
    if not params['program'] and not params['search']:
        messages.warning(request, 'Please select a program before exporting Excel.')
        return redirect('manage_students')

    group_filter = _normalize_group_filter(params['group'], params['program'])
    students = _filter_students_by_group(
        _attach_admissions(
            list(_get_students_queryset(params['search'], params['program'], params['verified']))
        ),
        group_filter,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Students'
    ws.views.sheetView[0].showGridLines = True

    headers = [
        'Registration No',
        'Full Name',
        "Father's Name",
        "Mother's Name",
        'Email',
        'Mobile',
        'Aadhaar',
        'Program',
        'Subject Group',
        'DSC Courses',
        'Course',
        'Selected Courses / Paper Name',
        'Application No',
        'Verified',
        'Registered On',
    ]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='082B49', end_color='082B49', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    data_font = Font(name='Calibri', size=10)
    data_align_left = Alignment(horizontal='left', vertical='center')
    data_align_center = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    for row_idx, s in enumerate(students, 2):
        admission = getattr(s, 'latest_admission', None)
        father = getattr(s, 'father_name', '') or (admission.father_name if admission else '')
        mother = getattr(s, 'mother_name', '') or (admission.mother_name if admission else '')

        row_data = [
            s.registration_no,
            s.full_name,
            father,
            mother,
            s.email or '',
            s.mobile or '',
            s.aadhaar or '',
            s.program_type,
            getattr(s, 'bsc_group_label', '') or '',
            getattr(s, 'dsc_courses_display', '') or '',
            s.course_name,
            _format_selected_courses_for_export(s),
            (admission.application_no if admission else '') or '',
            'Yes' if s.is_verified else 'No',
            s.created_date.strftime('%d-%m-%Y %H:%M') if s.created_date else '',
        ]

        ws.row_dimensions[row_idx].height = 22
        fill_to_apply = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            if fill_to_apply.fill_type:
                cell.fill = fill_to_apply

            if col_idx in (1, 6, 7, 13, 14, 15):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Students_{params['program'] or 'All'}.xlsx".replace(' ', '_').replace('/', '_')
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@admin_login_required
def admin_print_application(request, app_no):
    admission = get_object_or_404(StudentAdmission, application_no=app_no)
    context = get_print_context(admission)
    context['preview_mode'] = False
    context['admin_view'] = True
    return render(request, 'admissions/print_full.html', context)


@admin_login_required
def admin_download_pdf(request, app_no):
    from admissions.pdf import render_admission_pdf

    admission = get_object_or_404(StudentAdmission, application_no=app_no)
    pdf_bytes = render_admission_pdf(admission)
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{app_no}.pdf"'
    return response


@admin_login_required
@require_http_methods(['POST'])
def update_admission_status(request, pk):
    admission = StudentAdmission.objects.filter(pk=pk).first()
    if not admission:
        messages.error(request, 'Admission not found.')
        return redirect('admin_dashboard')
    new_status = request.POST.get('status')
    if new_status in ('Approved', 'Rejected', 'Pending', 'Submitted', 'Draft'):
        admission.status = new_status
        if new_status == 'Draft':
            admission.is_submitted = False
            admission.is_approved = False
            admission.save(update_fields=['status', 'is_submitted', 'is_approved'])
            student = Student.objects.filter(registration_no=admission.reg_no).first()
            if student:
                student.is_verified = False
                student.save(update_fields=['is_verified'])
            messages.success(request, 'Application status reset to Draft (Unverified, editable by student).')
        else:
            admission.save(update_fields=['status'])
            if new_status == 'Approved':
                student = Student.objects.filter(registration_no=admission.reg_no).first()
                if student:
                    student.is_verified = True
                    student.save(update_fields=['is_verified'])
            messages.success(request, f'Status updated to {new_status}.')
    return redirect(request.POST.get('next', 'admin_dashboard'))


@admin_login_required
@require_http_methods(['POST'])
def bulk_update_admission_status(request):
    valid_statuses = ('Approved', 'Rejected', 'Pending', 'Submitted', 'Draft')
    raw_ids = request.POST.getlist('admission_ids')
    admission_ids = []
    for raw_id in raw_ids:
        try:
            admission_ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue

    new_status = request.POST.get('status', '').strip()
    if not admission_ids:
        messages.error(request, 'Select at least one application.')
        return redirect('admin_dashboard')
    if new_status not in valid_statuses:
        messages.error(request, 'Choose a valid status.')
        return redirect('admin_dashboard')

    admissions = StudentAdmission.objects.filter(pk__in=admission_ids)
    if new_status == 'Draft':
        updated = 0
        for adm in admissions:
            adm.status = 'Draft'
            adm.is_submitted = False
            adm.is_approved = False
            adm.save(update_fields=['status', 'is_submitted', 'is_approved'])
            Student.objects.filter(registration_no=adm.reg_no).update(is_verified=False)
            updated += 1
    else:
        updated = admissions.update(status=new_status)
        if new_status == 'Approved':
            reg_nos = admissions.values_list('reg_no', flat=True)
            Student.objects.filter(registration_no__in=reg_nos).update(is_verified=True)

    if updated:
        messages.success(
            request,
            f'Updated {updated} application{"s" if updated != 1 else ""} to {new_status}.',
        )
    else:
        messages.error(request, 'No matching applications found.')
    return redirect('admin_dashboard')


@admin_login_required
@require_http_methods(['POST'])
def update_enrollment_status(request, pk):
    from admissions.models import StudentEnrollment
    from admissions.utils import generate_enrollment_number

    enrollment = get_object_or_404(StudentEnrollment, pk=pk)
    new_status = request.POST.get('status', '').strip()
    if new_status in ('Approved', 'Submitted', 'Draft'):
        enrollment.status = new_status
        if new_status == 'Approved':
            enrollment.is_submitted = True
            if not enrollment.enrollment_no:
                enrollment.enrollment_no = generate_enrollment_number()
            enrollment.save(update_fields=['status', 'is_submitted', 'enrollment_no'])
            Student.objects.filter(registration_no=enrollment.reg_no).update(is_verified=True)
            messages.success(request, f'Enrollment {enrollment.enrollment_no} for {enrollment.full_name} has been Approved / Accepted.')
        elif new_status == 'Draft':
            enrollment.is_submitted = False
            enrollment.save(update_fields=['status', 'is_submitted'])
            messages.success(request, f'Enrollment {enrollment.enrollment_no or enrollment.reg_no} for {enrollment.full_name} reset to Draft (student can now edit).')
        else:
            enrollment.save(update_fields=['status'])
            messages.success(request, f'Enrollment status updated to {new_status}.')
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or 'admin_dashboard'
    return redirect(next_url)


@admin_login_required
@require_http_methods(['POST'])
def bulk_update_enrollment_status(request):
    from admissions.models import StudentEnrollment
    from admissions.utils import generate_enrollment_number

    enrollment_ids = request.POST.getlist('enrollment_ids')
    new_status = request.POST.get('status', '').strip()
    valid_statuses = ('Approved', 'Draft')

    if not enrollment_ids:
        messages.error(request, 'Select at least one enrollment application.')
        return redirect(request.META.get('HTTP_REFERER') or 'admin_dashboard')
    if new_status not in valid_statuses:
        messages.error(request, 'Choose a valid status (Approve or Reset to Draft).')
        return redirect(request.META.get('HTTP_REFERER') or 'admin_dashboard')

    enrollments = StudentEnrollment.objects.filter(pk__in=enrollment_ids)
    count = 0
    if new_status == 'Approved':
        for enr in enrollments:
            enr.status = 'Approved'
            enr.is_submitted = True
            if not enr.enrollment_no:
                enr.enrollment_no = generate_enrollment_number()
            enr.save(update_fields=['status', 'is_submitted', 'enrollment_no'])
            Student.objects.filter(registration_no=enr.reg_no).update(is_verified=True)
            count += 1
        messages.success(request, f'Successfully Approved & Accepted {count} enrollment application(s).')
    elif new_status == 'Draft':
        for enr in enrollments:
            enr.status = 'Draft'
            enr.is_submitted = False
            enr.save(update_fields=['status', 'is_submitted'])
            count += 1
        messages.success(request, f'Reset {count} enrollment application(s) to Draft (students can now edit).')

    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or 'admin_dashboard'
    return redirect(next_url)


@admin_login_required
def manage_enrollments(request):
    from admissions.models import StudentEnrollment

    search = request.GET.get('search', '').strip()
    program_filter = request.GET.get('program', '').strip()
    if program_filter == 'ALL':
        program_filter = ''
    status_filter = request.GET.get('status', 'ALL').strip() or 'ALL'

    enrollments = StudentEnrollment.objects.all().select_related('student')

    if search:
        enrollments = enrollments.filter(
            Q(full_name__icontains=search)
            | Q(reg_no__icontains=search)
            | Q(enrollment_no__icontains=search)
            | Q(mobile__icontains=search)
            | Q(email__icontains=search)
            | Q(transaction_id__icontains=search)
        )

    if program_filter:
        enrollments = enrollments.filter(program_type=program_filter)

    if status_filter != 'ALL':
        enrollments = enrollments.filter(status=status_filter)

    total_count = StudentEnrollment.objects.count()
    submitted_count = StudentEnrollment.objects.filter(status='Submitted').count()
    approved_count = StudentEnrollment.objects.filter(status='Approved').count()
    draft_count = StudentEnrollment.objects.filter(status='Draft').count()

    program_types = get_program_names(active_only=False)

    return render(request, 'admin_panel/enrollments.html', {
        'enrollments': enrollments,
        'search': search,
        'program_filter': program_filter,
        'status_filter': status_filter,
        'program_types': program_types,
        'total_count': total_count,
        'submitted_count': submitted_count,
        'approved_count': approved_count,
        'draft_count': draft_count,
        'filtered_count': enrollments.count(),
    })


@admin_login_required
@require_http_methods(['GET', 'POST'])
def admin_edit_enrollment(request, pk):
    import base64
    import json
    from admissions.constants import MEDIUM_CHOICES, RELIGION_CHOICES
    from admissions.models import StudentEnrollment, StudentAdmission
    from admissions.utils import generate_enrollment_number
    from courses.models import Program, ProgramCourse
    from courses.subject_groups import get_bsc_subject_group_sections, is_bsc_program

    enrollment = get_object_or_404(StudentEnrollment.objects.select_related('student', 'admission'), pk=pk)
    student = enrollment.student
    admission = enrollment.admission or StudentAdmission.objects.filter(reg_no=enrollment.reg_no).first()

    # Determine available courses for the program
    program_type = enrollment.program_type or (admission.program_type if admission else '') or (student.program_type if student else '')
    courses = ProgramCourse.objects.filter(program_type__iexact=program_type).order_by('sort_order', 'course_code')
    if not courses.exists() and program_type:
        short_name = program_type.replace(' First Semester', '').strip()
        courses = ProgramCourse.objects.filter(program_type__iexact=short_name).order_by('sort_order', 'course_code')

    bsc_subject_groups = get_bsc_subject_group_sections(program_type) if is_bsc_program(program_type) else []

    if request.method == 'POST':
        # 1. Personal Details
        enrollment.full_name = request.POST.get('full_name', '').strip()
        enrollment.father_name = request.POST.get('father_name', '').strip()
        enrollment.mother_name = request.POST.get('mother_name', '').strip()
        enrollment.gender = request.POST.get('gender', '').strip()
        dob_str = request.POST.get('dob', '').strip()
        enrollment.dob = dob_str if dob_str else None
        enrollment.category = request.POST.get('category', '').strip()
        enrollment.nationality = request.POST.get('nationality', 'Indian').strip() or 'Indian'
        enrollment.religion = request.POST.get('religion', '').strip()
        enrollment.marital_status = request.POST.get('marital_status', '').strip()
        enrollment.blood_group = request.POST.get('blood_group', '').strip()
        enrollment.mobile = request.POST.get('mobile', '').strip()
        enrollment.email = request.POST.get('email', '').strip()
        enrollment.aadhaar = request.POST.get('aadhaar', '').strip()
        enrollment.apaar_id = request.POST.get('apaar_id', '').strip()
        enrollment.medium = request.POST.get('medium', '').strip()
        enrollment.has_disability = request.POST.get('has_disability') in ('1', 'true', 'True', True)
        enrollment.disability_details = request.POST.get('disability_details', '').strip()
        enrollment.disability_percentage = request.POST.get('disability_percentage', '').strip()
        enrollment.disability_type = request.POST.get('disability_type', '').strip()
        enrollment.is_minority = request.POST.get('is_minority') in ('1', 'true', 'True', True)

        # 2. Addresses
        enrollment.perm_state = request.POST.get('perm_state', '').strip()
        enrollment.perm_district = request.POST.get('perm_district', '').strip()
        enrollment.perm_city = request.POST.get('perm_city', '').strip()
        enrollment.perm_village = request.POST.get('perm_village', '').strip()
        enrollment.perm_pin_code = request.POST.get('perm_pin_code', '').strip()

        enrollment.corr_state = request.POST.get('corr_state', '').strip()
        enrollment.corr_district = request.POST.get('corr_district', '').strip()
        enrollment.corr_city = request.POST.get('corr_city', '').strip()
        enrollment.corr_village = request.POST.get('corr_village', '').strip()
        enrollment.corr_pin_code = request.POST.get('corr_pin_code', '').strip()

        # 3. Education Details
        enrollment.class10 = request.POST.get('class10', '10th').strip()
        enrollment.board10 = request.POST.get('board10', '').strip()
        y10 = request.POST.get('year10', '').strip()
        enrollment.year10 = int(y10) if y10.isdigit() else None
        enrollment.total_marks10 = request.POST.get('total_marks10', '').strip()
        enrollment.obtained10 = request.POST.get('obtained10', '').strip()
        enrollment.percentage10 = request.POST.get('percentage10', '').strip()
        enrollment.grade10 = request.POST.get('grade10', '').strip()

        enrollment.class12 = request.POST.get('class12', '12th').strip()
        enrollment.board12 = request.POST.get('board12', '').strip()
        enrollment.stream12 = request.POST.get('stream12', '').strip()
        y12 = request.POST.get('year12', '').strip()
        enrollment.year12 = int(y12) if y12.isdigit() else None
        enrollment.total_marks12 = request.POST.get('total_marks12', '').strip()
        enrollment.obtained12 = request.POST.get('obtained12', '').strip()
        enrollment.percentage12 = request.POST.get('percentage12', '').strip()
        enrollment.grade12 = request.POST.get('grade12', '').strip()

        enrollment.class_grad = request.POST.get('class_grad', '').strip()
        enrollment.board_grad = request.POST.get('board_grad', '').strip()
        enrollment.stream_grad = request.POST.get('stream_grad', '').strip()
        ygrad = request.POST.get('year_grad', '').strip()
        enrollment.year_grad = int(ygrad) if ygrad.isdigit() else None
        enrollment.total_marks_grad = request.POST.get('total_marks_grad', '').strip()
        enrollment.obtained_grad = request.POST.get('obtained_grad', '').strip()
        enrollment.percentage_grad = request.POST.get('percentage_grad', '').strip()
        enrollment.grade_grad = request.POST.get('grade_grad', '').strip()

        # 4. Fee & Payment Details
        enrollment.fee_amount = request.POST.get('fee_amount', '500').strip() or '500'
        enrollment.transaction_id = request.POST.get('transaction_id', '').strip()
        enrollment.payment_status = request.POST.get('payment_status', 'Paid').strip() or 'Paid'
        receipt_file = request.FILES.get('payment_receipt')
        if receipt_file:
            enrollment.payment_receipt = receipt_file

        # 5. Photo & Signature
        photo_file = request.FILES.get('photo_file')
        if photo_file:
            b64 = base64.b64encode(photo_file.read()).decode('utf-8')
            enrollment.photo_base64 = f"data:{photo_file.content_type};base64,{b64}"
        elif request.POST.get('photo_base64'):
            enrollment.photo_base64 = request.POST.get('photo_base64').strip()

        sig_file = request.FILES.get('signature_file')
        if sig_file:
            b64 = base64.b64encode(sig_file.read()).decode('utf-8')
            enrollment.signature_base64 = f"data:{sig_file.content_type};base64,{b64}"
        elif request.POST.get('signature_base64'):
            enrollment.signature_base64 = request.POST.get('signature_base64').strip()

        # 6. Courses Selection
        selected_course_ids = request.POST.getlist('selected_courses')
        bsc_subject_group = request.POST.get('bsc_subject_group', '').strip()
        raw_courses_json = request.POST.get('selected_courses_json', '').strip()

        if selected_course_ids:
            sel_courses = ProgramCourse.objects.filter(id__in=selected_course_ids).order_by('sort_order', 'course_code')
            course_list = []
            for c in sel_courses:
                course_list.append({
                    'id': str(c.id),
                    'code': c.course_code,
                    'name': c.course_name,
                    'paper': c.paper_no,
                    'type_1': c.course_type_1,
                    'type_2': c.course_type_2,
                    'dept': c.department,
                    'credit_l': c.credit_l,
                    'credit_t': c.credit_t,
                    'credit_p': c.credit_p,
                })
            if bsc_subject_group:
                enrollment.selected_courses_json = json.dumps({
                    'bsc_subject_group': bsc_subject_group,
                    'courses': course_list,
                    'subjects': course_list,
                })
            else:
                enrollment.selected_courses_json = json.dumps(course_list)
        elif raw_courses_json:
            enrollment.selected_courses_json = raw_courses_json

        # 7. Status, Remarks & Admin Action
        new_status = request.POST.get('status', 'Submitted').strip()
        admin_remarks = request.POST.get('admin_remarks', '').strip()
        enrollment.admin_remarks = admin_remarks

        if new_status == 'Approved':
            enrollment.status = 'Approved'
            enrollment.is_submitted = True
            if not enrollment.enrollment_no:
                enrollment.enrollment_no = generate_enrollment_number()
            student.is_verified = True
            student.save(update_fields=['is_verified'])
        elif new_status == 'Draft':
            enrollment.status = 'Draft'
            enrollment.is_submitted = False
        else:  # 'Submitted'
            enrollment.status = 'Submitted'
            enrollment.is_submitted = True
            if not enrollment.submitted_date:
                enrollment.submitted_date = timezone.now()

        enrollment.save()

        # Sync student basic fields for consistency
        if enrollment.full_name and student.full_name != enrollment.full_name:
            student.full_name = enrollment.full_name
        if enrollment.mobile and student.mobile != enrollment.mobile:
            student.mobile = enrollment.mobile
        if enrollment.email and student.email != enrollment.email:
            student.email = enrollment.email
        if enrollment.aadhaar and student.aadhaar != enrollment.aadhaar:
            student.aadhaar = enrollment.aadhaar
        student.save()

        status_msg = {
            'Approved': 'Approved / Accepted (locked for student)',
            'Submitted': 'Updated automatically at student end (Status: Submitted)',
            'Draft': 'Reset to Draft (Student must review & resubmit)',
        }.get(new_status, new_status)

        messages.success(
            request,
            f'Enrollment application #{enrollment.enrollment_no or enrollment.pk} for {enrollment.full_name} '
            f'successfully updated. Status: {status_msg}.'
        )

        if 'save_continue' in request.POST:
            return redirect('admin_edit_enrollment', pk=enrollment.pk)
        return redirect('manage_enrollments')

    # GET Request: Prepare form data
    selected_course_ids = []
    selected_course_codes = []
    current_bsc_group = ''
    if enrollment.selected_courses_json:
        try:
            data = json.loads(enrollment.selected_courses_json)
            c_list = []
            if isinstance(data, list):
                c_list = data
            elif isinstance(data, dict):
                c_list = data.get('courses') or data.get('subjects') or []
                current_bsc_group = data.get('bsc_subject_group') or data.get('BScSubjectGroup') or ''
            for c in c_list:
                if isinstance(c, dict):
                    if c.get('id'):
                        selected_course_ids.append(str(c.get('id')))
                    if c.get('code') or c.get('course_code'):
                        selected_course_codes.append((c.get('code') or c.get('course_code')).strip().upper())
        except Exception:
            pass

    # Annotate courses for convenient template checking
    annotated_courses = []
    for c in courses:
        is_sel = (str(c.id) in selected_course_ids) or (c.course_code.strip().upper() in selected_course_codes)
        annotated_courses.append({
            'obj': c,
            'is_selected': is_sel,
        })

    is_pg = any(x in (program_type or '').upper() for x in ('PG', 'M.SC', 'M.A', 'M.COM', 'M.S.W', 'PGDCA'))

    return render(request, 'admin_panel/edit_enrollment.html', {
        'enrollment': enrollment,
        'student': student,
        'admission': admission,
        'program_type': program_type,
        'courses': annotated_courses,
        'bsc_subject_groups': bsc_subject_groups,
        'current_bsc_group': current_bsc_group,
        'is_pg': is_pg,
        'religion_choices': RELIGION_CHOICES,
        'medium_choices': MEDIUM_CHOICES,
        'status_choices': StudentEnrollment.STATUS_CHOICES,
    })


def _format_enrollment_address(enrollment):
    perm_parts = [
        enrollment.perm_village,
        enrollment.perm_city,
        enrollment.perm_district,
        enrollment.perm_state,
        enrollment.perm_pin_code,
    ]
    parts = []
    for p in perm_parts:
        val = (p or '').strip()
        if val and (not parts or val.lower() != parts[-1].lower()):
            parts.append(val)

    if not parts:
        corr_parts = [
            enrollment.corr_village,
            enrollment.corr_city,
            enrollment.corr_district,
            enrollment.corr_state,
            enrollment.corr_pin_code,
        ]
        for p in corr_parts:
            val = (p or '').strip()
            if val and (not parts or val.lower() != parts[-1].lower()):
                parts.append(val)

    if not parts and enrollment.admission:
        adm = enrollment.admission
        adm_parts = [
            adm.perm_village,
            adm.perm_city,
            adm.perm_district,
            adm.perm_state,
            adm.perm_pin_code,
        ]
        for p in adm_parts:
            val = (p or '').strip()
            if val and (not parts or val.lower() != parts[-1].lower()):
                parts.append(val)

    return ', '.join(parts)


def _format_enrollment_gender(gender_val):
    if not gender_val:
        return ''
    g = str(gender_val).strip().lower()
    if g.startswith('m') or g in ('1', 'male'):
        return 'M'
    elif g.startswith('f') or g in ('0', 'female'):
        return 'F'
    return gender_val


def _format_enrollment_dob(dob_val):
    if not dob_val:
        return ''
    if hasattr(dob_val, 'strftime'):
        return dob_val.strftime('%m/%d/%Y')
    return str(dob_val)


def _format_enrollment_subject_codes(enrollment):
    import json
    from courses.models import ProgramCourse

    courses = []
    if enrollment.selected_courses_json:
        try:
            data = json.loads(enrollment.selected_courses_json)
            if isinstance(data, list):
                courses = data
            elif isinstance(data, dict):
                courses = data.get('courses') or data.get('subjects') or []
        except Exception:
            courses = []

    codes = []
    for c in courses:
        if not isinstance(c, dict):
            continue
        code = (c.get('code') or c.get('course_code') or '').strip()
        name = (c.get('name') or c.get('course_name') or '').strip()
        dept = (c.get('dept') or c.get('department') or '').strip()
        type_2 = (c.get('type_2') or c.get('type2') or '').strip()

        # If code was missing in JSON, resolve it from ProgramCourse
        if not code and name:
            clean_name = name.split('—')[-1].strip() if '—' in name else name
            db_c = ProgramCourse.objects.filter(program_type__iexact=enrollment.program_type).filter(
                Q(course_name__iexact=name) | Q(course_name__iexact=clean_name)
            ).first()
            if not db_c and dept:
                db_c = ProgramCourse.objects.filter(
                    program_type__iexact=enrollment.program_type,
                    department__iexact=dept,
                    course_type_2__iexact=type_2,
                ).first()
            if db_c and db_c.course_code:
                code = db_c.course_code.strip()

        if code and code not in codes:
            codes.append(code)

    return ', '.join(codes)


def _format_enrollment_subjects(enrollment):
    import json
    courses = []
    if enrollment.selected_courses_json:
        try:
            data = json.loads(enrollment.selected_courses_json)
            if isinstance(data, list):
                courses = data
            elif isinstance(data, dict):
                courses = data.get('courses') or data.get('subjects') or []
        except Exception:
            courses = []

    names = []
    for c in courses:
        raw_name = ''
        if isinstance(c, dict):
            raw_name = c.get('name') or c.get('course_name') or ''
        elif isinstance(c, str):
            raw_name = c
        if raw_name:
            clean = _export_course_name_only(raw_name)
            if clean and clean not in names:
                names.append(clean)

    if names:
        return ', '.join(names)

    adm = enrollment.admission
    if not adm and enrollment.reg_no:
        from admissions.models import StudentAdmission
        adm = StudentAdmission.objects.filter(reg_no=enrollment.reg_no).first()

    if adm and (adm.subject or '').strip():
        parts = [_export_course_name_only(p) for p in adm.subject.split(',') if _export_course_name_only(p)]
        return ', '.join(parts)

    return ''


def build_enrollment_export_row(enrollment):
    adm = enrollment.admission
    if not adm and enrollment.reg_no:
        from admissions.models import StudentAdmission
        adm = StudentAdmission.objects.filter(reg_no=enrollment.reg_no).first()

    admission_no = (getattr(adm, 'application_no', '') or enrollment.reg_no or '').strip()
    univ_enrol_no = enrollment.enrollment_no or ''
    stud_nm = enrollment.full_name or (enrollment.student.full_name if enrollment.student else '')
    father_name = enrollment.father_name or (getattr(adm, 'father_name', '') or '')
    mother_name = enrollment.mother_name or (getattr(adm, 'mother_name', '') or '')
    medium = enrollment.medium or (getattr(adm, 'medium', '') or '')
    category = enrollment.category or (getattr(adm, 'category', '') or '')

    gender_raw = enrollment.gender or (getattr(adm, 'gender', '') or '')
    gender_val = _format_enrollment_gender(gender_raw)

    dob_raw = enrollment.dob or (getattr(adm, 'dob', None))
    dob_val = _format_enrollment_dob(dob_raw)

    address = _format_enrollment_address(enrollment)
    mobile = enrollment.mobile or (enrollment.student.mobile if enrollment.student else '')
    class_name = enrollment.program_type or ''
    subject_codes = _format_enrollment_subject_codes(enrollment)
    subjects = _format_enrollment_subjects(enrollment)
    utr_no = (enrollment.transaction_id or getattr(adm, 'transaction_id', '') or '').strip()

    return [
        admission_no,
        univ_enrol_no,
        stud_nm,
        father_name,
        mother_name,
        medium,
        category,
        gender_val,
        dob_val,
        address,
        mobile,
        class_name,
        subject_codes,
        subjects,
        utr_no,
    ]


@admin_login_required
def export_enrollments_excel(request):
    """Export student enrollments as an Excel (.xlsx) spreadsheet matching university format."""
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from admissions.models import StudentEnrollment

    search = request.GET.get('search', '').strip()
    program_filter = request.GET.get('program', '').strip()
    if program_filter == 'ALL':
        program_filter = ''
    status_filter = request.GET.get('status', 'ALL').strip() or 'ALL'

    enrollments = StudentEnrollment.objects.all().select_related('student', 'admission')

    if search:
        enrollments = enrollments.filter(
            Q(full_name__icontains=search)
            | Q(reg_no__icontains=search)
            | Q(enrollment_no__icontains=search)
            | Q(mobile__icontains=search)
            | Q(email__icontains=search)
            | Q(transaction_id__icontains=search)
        )

    if program_filter:
        enrollments = enrollments.filter(program_type=program_filter)

    if status_filter != 'ALL':
        enrollments = enrollments.filter(status=status_filter)

    enrollments = enrollments.order_by('program_type', 'enrollment_no', 'full_name')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Enrollments'
    ws.views.sheetView[0].showGridLines = True

    headers = [
        'AddmissionNo',
        'Univ_EnrolNo',
        'Stud_nm',
        'FatherName',
        'MotherName',
        'Medium',
        'Category',
        'Gender',
        'DOB (MM/DD/YYYY)',
        'Address',
        'Mobile',
        'CLASS NAME',
        'SUBJECT CODE',
        'SUBJECTS',
        'UTR No',
    ]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='082B49')
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    row_num = 2
    for enr in enrollments:
        row_data = build_enrollment_export_row(enr)
        ws.row_dimensions[row_num].height = 20
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            if col_idx in (8, 9, 11, 15):  # Gender, DOB, Mobile, UTR No
                cell.alignment = center_align
            else:
                cell.alignment = left_align
        row_num += 1

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    if ws.column_dimensions.get('J'):
        ws.column_dimensions['J'].width = min(max(ws.column_dimensions['J'].width, 30), 50)
    if ws.column_dimensions.get('M'):
        ws.column_dimensions['M'].width = min(max(ws.column_dimensions['M'].width, 24), 45)
    if ws.column_dimensions.get('N'):
        ws.column_dimensions['N'].width = min(max(ws.column_dimensions['N'].width, 35), 65)
    if ws.column_dimensions.get('O'):
        ws.column_dimensions['O'].width = min(max(ws.column_dimensions['O'].width, 18), 32)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename_parts = ['enrollments']
    if program_filter:
        safe_prog = re.sub(r'[^\w\-]+', '_', program_filter).strip('_')
        filename_parts.append(safe_prog)
    if status_filter != 'ALL':
        filename_parts.append(status_filter.lower())
    filename_parts.append(timezone.now().strftime('%Y%m%d'))
    filename = f"{'_'.join(filename_parts)}.xlsx"

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# NEP UG ADMISSION CUM ENROLLMENT (SEMESTERS II, III, IV, V) ADMIN VIEWS
# ==============================================================================

@admin_login_required
def manage_nepug_enrollments(request):
    """Admin management page for NEP UG Admission cum Enrollment (2nd/3rd/4th/5th sem)."""
    from admissions.models import NepUgAdmissionEnrollment
    from admissions.nep_ug_views import NEP_UG_PROGRAM_CHOICES

    search = request.GET.get('search', '').strip()
    program_filter = request.GET.get('program', '').strip()
    if program_filter == 'ALL':
        program_filter = ''
    semester_filter = request.GET.get('semester', 'ALL').strip() or 'ALL'
    status_filter = request.GET.get('status', 'ALL').strip() or 'ALL'

    records = NepUgAdmissionEnrollment.objects.all().select_related('student')

    if search:
        records = records.filter(
            Q(full_name__icontains=search)
            | Q(reg_no__icontains=search)
            | Q(application_no__icontains=search)
            | Q(enrollment_no__icontains=search)
            | Q(previous_roll_no__icontains=search)
            | Q(previous_enrollment_no__icontains=search)
            | Q(mobile__icontains=search)
            | Q(email__icontains=search)
            | Q(transaction_id__icontains=search)
        )

    if program_filter:
        records = records.filter(program_type=program_filter)

    if semester_filter != 'ALL':
        records = records.filter(semester=semester_filter)

    if status_filter != 'ALL':
        records = records.filter(status=status_filter)

    total_count = NepUgAdmissionEnrollment.objects.count()
    submitted_count = NepUgAdmissionEnrollment.objects.filter(status='Submitted').count()
    approved_count = NepUgAdmissionEnrollment.objects.filter(status='Approved').count()
    draft_count = NepUgAdmissionEnrollment.objects.filter(status='Draft').count()

    sem_counts = {
        'II': NepUgAdmissionEnrollment.objects.filter(semester='II').count(),
        'III': NepUgAdmissionEnrollment.objects.filter(semester='III').count(),
        'IV': NepUgAdmissionEnrollment.objects.filter(semester='IV').count(),
        'V': NepUgAdmissionEnrollment.objects.filter(semester='V').count(),
    }

    program_choices = [c[0] for c in NEP_UG_PROGRAM_CHOICES]

    return render(request, 'admin_panel/nepug_enrollments.html', {
        'records': records,
        'search': search,
        'program_filter': program_filter,
        'semester_filter': semester_filter,
        'status_filter': status_filter,
        'program_choices': program_choices,
        'semester_choices': NepUgAdmissionEnrollment.SEMESTER_CHOICES,
        'total_count': total_count,
        'submitted_count': submitted_count,
        'approved_count': approved_count,
        'draft_count': draft_count,
        'sem_counts': sem_counts,
        'filtered_count': records.count(),
    })


@admin_login_required
@require_http_methods(['POST'])
def update_nepug_status(request, pk):
    from admissions.models import NepUgAdmissionEnrollment
    from admissions.utils import generate_enrollment_number

    record = get_object_or_404(NepUgAdmissionEnrollment, pk=pk)
    new_status = request.POST.get('status', '').strip()
    if new_status in ('Approved', 'Submitted', 'Draft'):
        record.status = new_status
        if new_status == 'Approved':
            record.is_submitted = True
            if not record.enrollment_no:
                record.enrollment_no = record.previous_enrollment_no or generate_enrollment_number()
            record.save(update_fields=['status', 'is_submitted', 'enrollment_no', 'updated_at'])
            Student.objects.filter(registration_no=record.reg_no).update(is_verified=True)
            messages.success(request, f'NEPUG application {record.application_no} for {record.full_name} has been Approved / Accepted.')
        elif new_status == 'Draft':
            record.is_submitted = False
            record.save(update_fields=['status', 'is_submitted', 'updated_at'])
            messages.success(request, f'NEPUG application {record.application_no} for {record.full_name} reset to Draft (student can now edit).')
        else:
            record.save(update_fields=['status', 'updated_at'])
            messages.success(request, f'NEPUG status updated to {new_status}.')

    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or 'manage_nepug_enrollments'
    return redirect(next_url)


@admin_login_required
@require_http_methods(['POST'])
def bulk_update_nepug_status(request):
    from admissions.models import NepUgAdmissionEnrollment
    from admissions.utils import generate_enrollment_number

    record_ids = request.POST.getlist('record_ids')
    new_status = request.POST.get('status', '').strip()
    valid_statuses = ('Approved', 'Draft')

    if not record_ids:
        messages.error(request, 'Select at least one NEPUG application.')
        return redirect(request.META.get('HTTP_REFERER') or 'manage_nepug_enrollments')
    if new_status not in valid_statuses:
        messages.error(request, 'Choose a valid status (Approve or Reset to Draft).')
        return redirect(request.META.get('HTTP_REFERER') or 'manage_nepug_enrollments')

    records = NepUgAdmissionEnrollment.objects.filter(pk__in=record_ids)
    count = 0
    if new_status == 'Approved':
        for rec in records:
            rec.status = 'Approved'
            rec.is_submitted = True
            if not rec.enrollment_no:
                rec.enrollment_no = rec.previous_enrollment_no or generate_enrollment_number()
            rec.save(update_fields=['status', 'is_submitted', 'enrollment_no', 'updated_at'])
            Student.objects.filter(registration_no=rec.reg_no).update(is_verified=True)
            count += 1
        messages.success(request, f'Successfully Approved & Accepted {count} NEPUG application(s).')
    elif new_status == 'Draft':
        for rec in records:
            rec.status = 'Draft'
            rec.is_submitted = False
            rec.save(update_fields=['status', 'is_submitted', 'updated_at'])
            count += 1
        messages.success(request, f'Reset {count} NEPUG application(s) to Draft (students can now edit).')

    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or 'manage_nepug_enrollments'
    return redirect(next_url)


@admin_login_required
@require_http_methods(['GET', 'POST'])
def admin_edit_nepug(request, pk):
    import base64
    import json
    from datetime import datetime
    from accounts.models import Student
    from admissions.constants import MEDIUM_CHOICES, RELIGION_CHOICES
    from admissions.models import NepUgAdmissionEnrollment
    from admissions.nep_ug_views import NEP_UG_PROGRAM_CHOICES

    record = get_object_or_404(NepUgAdmissionEnrollment, pk=pk)

    if request.method == 'POST':
        # Update program & semester
        record.program_type = request.POST.get('program_type', '').strip() or record.program_type
        record.semester = request.POST.get('semester', '').strip() or record.semester
        record.academic_session = request.POST.get('academic_session', '').strip() or record.academic_session
        record.enrollment_no = request.POST.get('enrollment_no', '').strip() or record.enrollment_no
        record.previous_roll_no = request.POST.get('previous_roll_no', '').strip()
        record.previous_enrollment_no = request.POST.get('previous_enrollment_no', '').strip()
        record.previous_semester_result = request.POST.get('previous_semester_result', '').strip()
        record.previous_semester_marks = request.POST.get('previous_semester_marks', '').strip()
        record.previous_exam_name = request.POST.get('previous_exam_name', '').strip()
        record.previous_board = request.POST.get('previous_board', 'Shaheed Nandkumar Patel Vishwavidyalaya, Raigarh').strip()
        py_str = request.POST.get('previous_year', '').strip()
        record.previous_year = int(py_str) if py_str.isdigit() else None
        record.previous_total_marks = request.POST.get('previous_total_marks', '').strip()
        record.previous_obtained_marks = request.POST.get('previous_obtained_marks', '').strip()
        record.previous_percentage = request.POST.get('previous_percentage', '').strip()
        if not record.previous_semester_marks and record.previous_obtained_marks:
            record.previous_semester_marks = record.previous_obtained_marks

        # Update personal & contact
        record.full_name = request.POST.get('full_name', '').strip()
        record.father_name = request.POST.get('father_name', '').strip()
        record.mother_name = request.POST.get('mother_name', '').strip()
        record.gender = request.POST.get('gender', '').strip()
        record.category = request.POST.get('category', '').strip()
        record.religion = request.POST.get('religion', '').strip()
        record.blood_group = request.POST.get('blood_group', '').strip()
        record.medium = request.POST.get('medium', '').strip()
        record.mobile = request.POST.get('mobile', '').strip()
        record.email = request.POST.get('email', '').strip()
        record.aadhaar = request.POST.get('aadhaar', '').strip()
        record.apaar_id = request.POST.get('apaar_id', '').strip()

        dob_str = request.POST.get('dob', '').strip()
        if dob_str:
            for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                try:
                    record.dob = datetime.strptime(dob_str, fmt).date()
                    break
                except ValueError:
                    pass

        # Address
        record.corr_village = request.POST.get('corr_village', '').strip()
        record.corr_city = request.POST.get('corr_city', '').strip()
        record.corr_district = request.POST.get('corr_district', '').strip()
        record.corr_state = request.POST.get('corr_state', '').strip()
        record.corr_pin_code = request.POST.get('corr_pin_code', '').strip()

        # Education 10th
        record.class10 = request.POST.get('class10', '10th').strip() or '10th'
        record.board10 = request.POST.get('board10', '').strip()
        y10_str = request.POST.get('year10', '').strip()
        record.year10 = int(y10_str) if y10_str.isdigit() else None
        record.total_marks10 = request.POST.get('total_marks10', '').strip()
        record.obtained10 = request.POST.get('obtained10', '').strip()
        record.percentage10 = request.POST.get('percentage10', '').strip()
        record.grade10 = request.POST.get('grade10', '').strip()

        # Education 12th
        record.class12 = request.POST.get('class12', '12th').strip() or '12th'
        record.board12 = request.POST.get('board12', '').strip()
        record.stream12 = request.POST.get('stream12', '').strip()
        y12_str = request.POST.get('year12', '').strip()
        record.year12 = int(y12_str) if y12_str.isdigit() else None
        record.total_marks12 = request.POST.get('total_marks12', '').strip()
        record.obtained12 = request.POST.get('obtained12', '').strip()
        record.percentage12 = request.POST.get('percentage12', '').strip()
        record.grade12 = request.POST.get('grade12', '').strip()

        # Photo & Signature Upload
        if 'photo_file' in request.FILES and request.FILES['photo_file']:
            pf = request.FILES['photo_file']
            ctype = pf.content_type or 'image/jpeg'
            record.photo_base64 = f"data:{ctype};base64,{base64.b64encode(pf.read()).decode('utf-8')}"
        elif request.POST.get('photo_base64'):
            record.photo_base64 = request.POST.get('photo_base64').strip()

        if 'signature_file' in request.FILES and request.FILES['signature_file']:
            sf = request.FILES['signature_file']
            ctype = sf.content_type or 'image/jpeg'
            record.signature_base64 = f"data:{ctype};base64,{base64.b64encode(sf.read()).decode('utf-8')}"
        elif request.POST.get('signature_base64'):
            record.signature_base64 = request.POST.get('signature_base64').strip()

        # Education JSON
        education_list = []
        if record.class10 or record.board10:
            education_list.append({
                'RowKey': '10', 'className': record.class10, 'board': record.board10,
                'year': record.year10, 'totalMarks': record.total_marks10,
                'obtained': record.obtained10, 'percentage': record.percentage10, 'grade': record.grade10,
            })
        if record.class12 or record.board12:
            education_list.append({
                'RowKey': '12', 'className': record.class12, 'board': record.board12,
                'stream': record.stream12, 'year': record.year12, 'totalMarks': record.total_marks12,
                'obtained': record.obtained12, 'percentage': record.percentage12, 'grade': record.grade12,
            })
        if record.previous_exam_name or record.previous_roll_no:
            education_list.append({
                'RowKey': 'ug', 'className': record.previous_exam_name or f"Semester Exam",
                'board': record.previous_board, 'rollNo': record.previous_roll_no,
                'enrollmentNo': record.previous_enrollment_no, 'year': record.previous_year,
                'totalMarks': record.previous_total_marks,
                'obtained': record.previous_obtained_marks or record.previous_semester_marks,
                'percentage': record.previous_percentage, 'result': record.previous_semester_result,
            })
        record.education_json = json.dumps(education_list)

        # Payment & Status
        record.fee_amount = request.POST.get('fee_amount', '500').strip()
        record.transaction_id = request.POST.get('transaction_id', '').strip()
        record.payment_status = request.POST.get('payment_status', 'Paid').strip()
        record.admin_remarks = request.POST.get('admin_remarks', '').strip()

        new_status = request.POST.get('status', '').strip()
        if new_status in ('Approved', 'Submitted', 'Draft'):
            record.status = new_status
            if new_status == 'Approved':
                record.is_submitted = True
                Student.objects.filter(registration_no=record.reg_no).update(
                    is_verified=True,
                    full_name=record.full_name,
                    mobile=record.mobile,
                    email=record.email,
                )
            elif new_status == 'Draft':
                record.is_submitted = False

        record.save()
        messages.success(request, f'NEPUG record for {record.full_name} ({record.application_no}) updated successfully.')
        return redirect('manage_nepug_enrollments')

    return render(request, 'admin_panel/edit_nepug.html', {
        'record': record,
        'nep_program_choices': NEP_UG_PROGRAM_CHOICES,
        'semester_choices': NepUgAdmissionEnrollment.SEMESTER_CHOICES,
        'religion_choices': RELIGION_CHOICES,
        'medium_choices': MEDIUM_CHOICES,
        'status_choices': NepUgAdmissionEnrollment.STATUS_CHOICES,
    })


@admin_login_required
@require_http_methods(['GET', 'POST'])
def admin_create_nepug(request):
    """Admin entry creation for ex-students continuing into NEP UG (2nd, 3rd, 4th, 5th sem)."""
    import base64
    import json
    from datetime import datetime
    from accounts.models import Student
    from admissions.constants import MEDIUM_CHOICES, RELIGION_CHOICES
    from admissions.models import NepUgAdmissionEnrollment, StudentAdmission, StudentEnrollment
    from admissions.nep_ug_views import NEP_UG_PROGRAM_CHOICES, generate_nep_ug_application_number
    from admissions.utils import generate_enrollment_number

    initial_data = {}
    lookup_reg = request.GET.get('reg_no', '').strip()
    if lookup_reg:
        std = Student.objects.filter(registration_no=lookup_reg).first()
        adm = StudentAdmission.objects.filter(reg_no=lookup_reg).order_by('-submitted_date', '-created_date').first()
        enr = StudentEnrollment.objects.filter(reg_no=lookup_reg).order_by('-submitted_date', '-created_at').first()
        nep = NepUgAdmissionEnrollment.objects.filter(reg_no=lookup_reg).order_by('-submitted_date', '-created_at').first()

        if std:
            initial_data['reg_no'] = std.registration_no
            initial_data['full_name'] = std.full_name
            initial_data['email'] = std.email
            initial_data['mobile'] = std.mobile
            initial_data['aadhaar'] = std.aadhaar
            if std.program_type:
                initial_data['program_type'] = std.program_type.replace(' - First Semester', '').replace(' First Semester', '').strip()

        if adm:
            initial_data['father_name'] = adm.father_name
            initial_data['mother_name'] = adm.mother_name
            initial_data['gender'] = adm.gender
            initial_data['category'] = adm.category
            initial_data['dob'] = adm.dob.strftime('%Y-%m-%d') if adm.dob else ''
            initial_data['medium'] = adm.medium
            initial_data['corr_village'] = adm.corr_village or adm.perm_village
            initial_data['corr_city'] = adm.corr_city or adm.perm_city
            initial_data['corr_district'] = adm.corr_district or adm.perm_district
            initial_data['corr_state'] = adm.corr_state or adm.perm_state
            initial_data['corr_pin_code'] = adm.corr_pin_code or adm.perm_pin_code
            if not initial_data.get('program_type') and adm.program_type:
                initial_data['program_type'] = adm.program_type.replace(' - First Semester', '').replace(' First Semester', '').strip()

        if enr:
            initial_data['previous_enrollment_no'] = enr.enrollment_no
            if not initial_data.get('program_type') and enr.program_type:
                initial_data['program_type'] = enr.program_type.replace(' - First Semester', '').replace(' First Semester', '').strip()
            if not initial_data.get('father_name') and enr.father_name:
                initial_data['father_name'] = enr.father_name
            if not initial_data.get('mother_name') and enr.mother_name:
                initial_data['mother_name'] = enr.mother_name

        # Helper to get first non-empty value across nep, enr, adm
        def _get_val(attr, default=''):
            for obj in (nep, enr, adm):
                if obj and hasattr(obj, attr):
                    val = getattr(obj, attr)
                    if val is not None and str(val).strip() != '':
                        return str(val)
            return default

        # 10th details prefill
        initial_data['class10'] = _get_val('class10', '10th')
        initial_data['board10'] = _get_val('board10', 'CGBSE')
        initial_data['year10'] = _get_val('year10', '')
        initial_data['total_marks10'] = _get_val('total_marks10', '')
        initial_data['obtained10'] = _get_val('obtained10', '')
        initial_data['percentage10'] = _get_val('percentage10', '')
        initial_data['grade10'] = _get_val('grade10', '')

        # 12th details prefill
        initial_data['class12'] = _get_val('class12', '12th')
        initial_data['board12'] = _get_val('board12', 'CGBSE')
        initial_data['stream12'] = _get_val('stream12', '')
        initial_data['year12'] = _get_val('year12', '')
        initial_data['total_marks12'] = _get_val('total_marks12', '')
        initial_data['obtained12'] = _get_val('obtained12', '')
        initial_data['percentage12'] = _get_val('percentage12', '')
        initial_data['grade12'] = _get_val('grade12', '')

        # Previous UG Marksheet / Semester details prefill
        if nep:
            initial_data['previous_roll_no'] = nep.previous_roll_no
            initial_data['previous_enrollment_no'] = nep.previous_enrollment_no or nep.enrollment_no
            initial_data['previous_exam_name'] = nep.previous_exam_name or f"{nep.program_type} Semester {nep.semester}"
            initial_data['previous_board'] = nep.previous_board or 'Shaheed Nandkumar Patel Vishwavidyalaya, Raigarh'
            initial_data['previous_year'] = nep.previous_year or ''
            initial_data['previous_total_marks'] = nep.previous_total_marks or ''
            initial_data['previous_obtained_marks'] = nep.previous_obtained_marks or nep.previous_semester_marks or ''
            initial_data['previous_percentage'] = nep.previous_percentage or ''
            initial_data['previous_semester_result'] = nep.previous_semester_result or 'Pass'
            initial_data['previous_semester_marks'] = nep.previous_semester_marks or ''
        elif enr:
            initial_data['previous_enrollment_no'] = enr.enrollment_no
            initial_data['previous_exam_name'] = f"{enr.program_type} Semester {enr.semester}"
            initial_data['previous_board'] = 'Shaheed Nandkumar Patel Vishwavidyalaya, Raigarh'
            initial_data['previous_year'] = timezone.now().year
        elif adm:
            initial_data['previous_exam_name'] = f"{adm.program_type} Semester I"
            initial_data['previous_board'] = 'Shaheed Nandkumar Patel Vishwavidyalaya, Raigarh'
            initial_data['previous_year'] = timezone.now().year

        # Photo & Signature prefill
        photo_src = (nep.photo_base64 if nep and nep.photo_base64 else None) or \
                    (enr.photo_base64 if enr and enr.photo_base64 else None) or \
                    (adm.photo_base64 if adm and adm.photo_base64 else None)
        if photo_src:
            initial_data['photo_base64'] = photo_src

        sig_src = (nep.signature_base64 if nep and nep.signature_base64 else None) or \
                  (enr.signature_base64 if enr and enr.signature_base64 else None) or \
                  (adm.signature_base64 if adm and adm.signature_base64 else None)
        if sig_src:
            initial_data['signature_base64'] = sig_src

    if request.method == 'POST':
        reg_no = request.POST.get('reg_no', '').strip()
        full_name = request.POST.get('full_name', '').strip()
        if not reg_no or not full_name:
            messages.error(request, 'Registration No and Student Full Name are required.')
            return render(request, 'admin_panel/create_nepug.html', {
                'initial_data': request.POST.dict(),
                'nep_program_choices': NEP_UG_PROGRAM_CHOICES,
                'semester_choices': NepUgAdmissionEnrollment.SEMESTER_CHOICES,
                'religion_choices': RELIGION_CHOICES,
                'medium_choices': MEDIUM_CHOICES,
                'status_choices': NepUgAdmissionEnrollment.STATUS_CHOICES,
            })

        student = Student.objects.filter(registration_no=reg_no).first()
        if not student:
            student = Student.objects.create(
                registration_no=reg_no,
                full_name=full_name,
                email=request.POST.get('email', '').strip() or f'{reg_no.lower()}@college.local',
                mobile=request.POST.get('mobile', '').strip() or '0000000000',
                password='password123',
                is_verified=True,
            )

        adm = StudentAdmission.objects.filter(reg_no=reg_no).first()
        first_enr = StudentEnrollment.objects.filter(reg_no=reg_no).first()

        semester = request.POST.get('semester', 'II').strip() or 'II'
        program_type = request.POST.get('program_type', 'B.A.').strip()
        academic_session = request.POST.get('academic_session', '2026-27').strip() or '2026-27'
        previous_roll_no = request.POST.get('previous_roll_no', '').strip()
        previous_enrollment_no = request.POST.get('previous_enrollment_no', '').strip()
        previous_semester_result = request.POST.get('previous_semester_result', 'Pass').strip()
        previous_semester_marks = request.POST.get('previous_semester_marks', '').strip()
        previous_exam_name = request.POST.get('previous_exam_name', '').strip()
        previous_board = request.POST.get('previous_board', 'Shaheed Nandkumar Patel Vishwavidyalaya, Raigarh').strip()
        py_str = request.POST.get('previous_year', '').strip()
        previous_year = int(py_str) if py_str.isdigit() else None
        previous_total_marks = request.POST.get('previous_total_marks', '').strip()
        previous_obtained_marks = request.POST.get('previous_obtained_marks', '').strip()
        previous_percentage = request.POST.get('previous_percentage', '').strip()
        if not previous_semester_marks and previous_obtained_marks:
            previous_semester_marks = previous_obtained_marks

        # 10th details
        class10 = request.POST.get('class10', '10th').strip() or '10th'
        board10 = request.POST.get('board10', '').strip()
        y10_str = request.POST.get('year10', '').strip()
        year10 = int(y10_str) if y10_str.isdigit() else None
        total_marks10 = request.POST.get('total_marks10', '').strip()
        obtained10 = request.POST.get('obtained10', '').strip()
        percentage10 = request.POST.get('percentage10', '').strip()
        grade10 = request.POST.get('grade10', '').strip()

        # 12th details
        class12 = request.POST.get('class12', '12th').strip() or '12th'
        board12 = request.POST.get('board12', '').strip()
        stream12 = request.POST.get('stream12', '').strip()
        y12_str = request.POST.get('year12', '').strip()
        year12 = int(y12_str) if y12_str.isdigit() else None
        total_marks12 = request.POST.get('total_marks12', '').strip()
        obtained12 = request.POST.get('obtained12', '').strip()
        percentage12 = request.POST.get('percentage12', '').strip()
        grade12 = request.POST.get('grade12', '').strip()

        # Handle Photo Upload
        photo_base64 = request.POST.get('photo_base64', '').strip()
        if 'photo_file' in request.FILES and request.FILES['photo_file']:
            pf = request.FILES['photo_file']
            ctype = pf.content_type or 'image/jpeg'
            photo_base64 = f"data:{ctype};base64,{base64.b64encode(pf.read()).decode('utf-8')}"

        # Handle Signature Upload
        signature_base64 = request.POST.get('signature_base64', '').strip()
        if 'signature_file' in request.FILES and request.FILES['signature_file']:
            sf = request.FILES['signature_file']
            ctype = sf.content_type or 'image/jpeg'
            signature_base64 = f"data:{ctype};base64,{base64.b64encode(sf.read()).decode('utf-8')}"

        # Build education_json
        education_list = []
        if class10 or board10:
            education_list.append({
                'RowKey': '10', 'className': class10, 'board': board10,
                'year': year10, 'totalMarks': total_marks10,
                'obtained': obtained10, 'percentage': percentage10, 'grade': grade10,
            })
        if class12 or board12:
            education_list.append({
                'RowKey': '12', 'className': class12, 'board': board12,
                'stream': stream12, 'year': year12, 'totalMarks': total_marks12,
                'obtained': obtained12, 'percentage': percentage12, 'grade': grade12,
            })
        if previous_exam_name or previous_roll_no:
            education_list.append({
                'RowKey': 'ug', 'className': previous_exam_name or f"Semester Exam",
                'board': previous_board, 'rollNo': previous_roll_no,
                'enrollmentNo': previous_enrollment_no, 'year': previous_year,
                'totalMarks': previous_total_marks,
                'obtained': previous_obtained_marks or previous_semester_marks,
                'percentage': previous_percentage, 'result': previous_semester_result,
            })

        dob_val = None
        dob_str = request.POST.get('dob', '').strip()
        if dob_str:
            for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                try:
                    dob_val = datetime.strptime(dob_str, fmt).date()
                    break
                except ValueError:
                    pass

        status = request.POST.get('status', 'Approved').strip() or 'Approved'
        fee_amount = request.POST.get('fee_amount', '500').strip() or '500'
        transaction_id = request.POST.get('transaction_id', '').strip()
        payment_status = request.POST.get('payment_status', 'Paid').strip() or 'Paid'

        app_no = generate_nep_ug_application_number()
        enrol_no = previous_enrollment_no or generate_enrollment_number()

        rec = NepUgAdmissionEnrollment.objects.create(
            application_no=app_no,
            reg_no=reg_no,
            enrollment_no=enrol_no,
            student=student,
            admission=adm,
            first_sem_enrollment=first_enr,
            program_type=program_type,
            semester=semester,
            academic_session=academic_session,
            previous_roll_no=previous_roll_no,
            previous_enrollment_no=previous_enrollment_no,
            previous_semester_result=previous_semester_result,
            previous_semester_marks=previous_semester_marks,
            previous_exam_name=previous_exam_name,
            previous_board=previous_board,
            previous_year=previous_year,
            previous_total_marks=previous_total_marks,
            previous_obtained_marks=previous_obtained_marks,
            previous_percentage=previous_percentage,
            class10=class10,
            board10=board10,
            year10=year10,
            total_marks10=total_marks10,
            obtained10=obtained10,
            percentage10=percentage10,
            grade10=grade10,
            class12=class12,
            board12=board12,
            stream12=stream12,
            year12=year12,
            total_marks12=total_marks12,
            obtained12=obtained12,
            percentage12=percentage12,
            grade12=grade12,
            education_json=json.dumps(education_list),
            photo_base64=photo_base64,
            signature_base64=signature_base64,
            full_name=full_name,
            father_name=request.POST.get('father_name', '').strip(),
            mother_name=request.POST.get('mother_name', '').strip(),
            gender=request.POST.get('gender', '').strip(),
            dob=dob_val,
            category=request.POST.get('category', '').strip(),
            nationality=request.POST.get('nationality', 'Indian').strip(),
            religion=request.POST.get('religion', '').strip(),
            blood_group=request.POST.get('blood_group', '').strip(),
            mobile=request.POST.get('mobile', '').strip() or student.mobile,
            email=request.POST.get('email', '').strip() or student.email,
            aadhaar=request.POST.get('aadhaar', '').strip(),
            apaar_id=request.POST.get('apaar_id', '').strip(),
            medium=request.POST.get('medium', 'Hindi').strip(),
            corr_village=request.POST.get('corr_village', '').strip(),
            corr_city=request.POST.get('corr_city', '').strip(),
            corr_district=request.POST.get('corr_district', '').strip(),
            corr_state=request.POST.get('corr_state', 'Chhattisgarh').strip(),
            corr_pin_code=request.POST.get('corr_pin_code', '').strip(),
            fee_amount=fee_amount,
            transaction_id=transaction_id,
            payment_status=payment_status,
            status=status,
            is_submitted=(status in ('Approved', 'Submitted')),
            submitted_date=timezone.now(),
            admin_remarks=request.POST.get('admin_remarks', '').strip(),
        )

        # Sync Student model
        student.is_verified = True
        student.full_name = full_name
        student.mobile = rec.mobile
        student.email = rec.email
        student.aadhaar = rec.aadhaar
        if not student.program_type:
            student.program_type = program_type
        student.save()

        messages.success(request, f'NEPUG record created successfully for {rec.full_name} ({rec.application_no})!')
        return redirect('manage_nepug_enrollments')

    return render(request, 'admin_panel/create_nepug.html', {
        'initial_data': initial_data,
        'nep_program_choices': NEP_UG_PROGRAM_CHOICES,
        'semester_choices': NepUgAdmissionEnrollment.SEMESTER_CHOICES,
        'religion_choices': RELIGION_CHOICES,
        'medium_choices': MEDIUM_CHOICES,
        'status_choices': NepUgAdmissionEnrollment.STATUS_CHOICES,
    })


@admin_login_required
def admin_print_nepug(request, pk):
    from admissions.models import NepUgAdmissionEnrollment
    nep_record = get_object_or_404(NepUgAdmissionEnrollment, pk=pk)
    return render(request, 'admissions/nep_ug_print.html', {
        'student': nep_record.student,
        'nep': nep_record,
    })


@admin_login_required
def admin_receipt_nepug(request, pk):
    from admissions.models import NepUgAdmissionEnrollment
    nep_record = get_object_or_404(NepUgAdmissionEnrollment, pk=pk)
    sub_date = nep_record.submitted_date or timezone.now()
    deposit_date = sub_date.strftime('%d/%m/%Y at %I:%M %p')
    month_year = sub_date.strftime('%b. / %Y').upper()
    program_display = f"{nep_record.program_type} ({nep_record.semester_display})"

    return render(request, 'admissions/nep_ug_receipt.html', {
        'student': nep_record.student,
        'nep': nep_record,
        'program_display': program_display,
        'deposit_date': deposit_date,
        'month_year': month_year,
        'fee_amount': nep_record.fee_amount or '500',
        'transaction_id': nep_record.transaction_id or '',
        'payment_status': nep_record.payment_status or ('Paid' if nep_record.transaction_id else 'Pending'),
    })


@admin_login_required
def export_nepug_excel(request):
    """Export NEPUG applications to Excel (.xlsx) matching university format."""
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from admissions.models import NepUgAdmissionEnrollment

    search = request.GET.get('search', '').strip()
    program_filter = request.GET.get('program', '').strip()
    if program_filter == 'ALL':
        program_filter = ''
    semester_filter = request.GET.get('semester', 'ALL').strip() or 'ALL'
    status_filter = request.GET.get('status', 'ALL').strip() or 'ALL'

    records = NepUgAdmissionEnrollment.objects.all().select_related('student')

    if search:
        records = records.filter(
            Q(full_name__icontains=search)
            | Q(reg_no__icontains=search)
            | Q(application_no__icontains=search)
            | Q(enrollment_no__icontains=search)
            | Q(previous_roll_no__icontains=search)
            | Q(previous_enrollment_no__icontains=search)
            | Q(mobile__icontains=search)
            | Q(email__icontains=search)
            | Q(transaction_id__icontains=search)
        )

    if program_filter:
        records = records.filter(program_type=program_filter)

    if semester_filter != 'ALL':
        records = records.filter(semester=semester_filter)

    if status_filter != 'ALL':
        records = records.filter(status=status_filter)

    records = records.order_by('semester', 'program_type', 'application_no')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'NEPUG_Enrollments'
    ws.views.sheetView[0].showGridLines = True

    headers = [
        'ApplicationNo',
        'Univ_EnrolNo',
        'Semester',
        'Stud_nm',
        'FatherName',
        'MotherName',
        'Medium',
        'Category',
        'Gender',
        'DOB (MM/DD/YYYY)',
        'Address',
        'Mobile',
        'Email',
        'CLASS NAME',
        'Prev Roll No',
        'Prev Result',
        'Prev Marks',
        'UTR No',
        'Fee Amount',
        'Status',
        'Submitted Date',
    ]

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='082B49')
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    ws.row_dimensions[1].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    center_col_indices = {1, 2, 3, 7, 8, 9, 10, 12, 15, 16, 18, 19, 20, 21}

    for row_idx, rec in enumerate(records, start=2):
        gender_val = _format_enrollment_gender(rec.gender)
        dob_val = _format_enrollment_dob(rec.dob)
        addr = _format_enrollment_address(rec)
        sub_date = rec.submitted_date.strftime('%d/%m/%Y %I:%M %p') if rec.submitted_date else ''

        row_data = [
            rec.application_no or '',
            rec.enrollment_no or rec.previous_enrollment_no or '',
            rec.semester or '',
            rec.full_name or '',
            rec.father_name or '',
            rec.mother_name or '',
            rec.medium or '',
            rec.category or '',
            gender_val,
            dob_val,
            addr,
            rec.mobile or '',
            rec.email or '',
            rec.program_type or '',
            rec.previous_roll_no or '',
            rec.previous_semester_result or '',
            rec.previous_semester_marks or '',
            rec.transaction_id or '',
            rec.fee_amount or '500',
            rec.status or '',
            sub_date,
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name='Calibri', size=10)
            cell.border = thin_border
            if col_idx in center_col_indices:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"nepug_enrollments_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


ATTENDANCE_MONTHS = [
    'January', 'February', 'March', 'April', 'May', 'June',
    'July', 'August', 'September', 'October', 'November', 'December'
]


def _normalize_attendance_str(text):
    if not text:
        return ''
    return re.sub(r'[^\w]+', '', str(text).lower())


def _course_item_matches(item, selected_course):
    """Check if a subject dict from admission or enrollment JSON matches the ProgramCourse."""
    if not isinstance(item, dict):
        return False

    # 1. Direct ID match
    item_id = str(item.get('id') or '').strip()
    if item_id and item_id == str(selected_course.id):
        return True

    # 2. Course code match
    selected_code = _normalize_attendance_str(selected_course.course_code)
    item_code = _normalize_attendance_str(item.get('code') or item.get('course_code') or '')
    if selected_code and item_code and selected_code == item_code:
        return True

    # 3. Type check (Theory vs Practical)
    c_type1 = (selected_course.course_type_1 or '').strip().lower()
    item_type1 = (item.get('type1') or item.get('type_1') or '').strip().lower()
    if c_type1 and item_type1:
        is_c_pract = ('pract' in c_type1 or 'lab' in c_type1)
        is_it_pract = ('pract' in item_type1 or 'lab' in item_type1)
        if is_c_pract != is_it_pract:
            return False

    # 4. Name and Department match
    item_name = (item.get('name') or item.get('course_name') or '').strip()
    item_dept = (item.get('dept') or item.get('department') or '').strip()
    c_name = (selected_course.course_name or '').strip()
    c_dept = (selected_course.department or '').strip()

    # Split department prefix in item_name if present (e.g. "Geography — Fundamental of Physical Geography")
    if ' — ' in item_name:
        parts = item_name.split(' — ', 1)
        item_dept = item_dept or parts[0].strip()
        item_cname = parts[1].strip()
    elif ' - ' in item_name:
        parts = item_name.split(' - ', 1)
        item_dept = item_dept or parts[0].strip()
        item_cname = parts[1].strip()
    else:
        item_cname = item_name

    norm_c_name = _normalize_attendance_str(c_name)
    norm_it_cname = _normalize_attendance_str(item_cname)
    norm_it_full = _normalize_attendance_str(item_name)

    if norm_c_name and (norm_c_name in norm_it_full or norm_it_cname in norm_c_name or norm_c_name in norm_it_cname):
        return True

    # Hindi script matching support
    if c_dept.lower() == 'hindi' and (item_dept.lower() == 'hindi' or 'hindi' in item_name.lower()):
        if 'sahitya' in norm_c_name and ('साहित्य' in item_name or 'sahitya' in norm_it_full):
            return True
        if 'language' in norm_c_name and ('language' in norm_it_full or 'भाषा' in item_name):
            return True

    return False


def _student_takes_course(student, selected_course):
    """Determine whether student took selected_course from their enrollment or admission."""
    # Check enrollment first
    enr = getattr(student, 'latest_enrollment', None)
    if enr and enr.selected_courses_json:
        try:
            data = json.loads(enr.selected_courses_json)
            c_list = data if isinstance(data, list) else (data.get('courses') or data.get('subjects') or [])
            for it in c_list:
                if _course_item_matches(it, selected_course):
                    return True
        except Exception:
            pass

    # Check admission
    adm = getattr(student, 'latest_admission', None)
    if adm and adm.selected_subjects_json:
        try:
            items, _ = parse_selected_subjects_payload(adm.selected_subjects_json)
            for it in items:
                if _course_item_matches(it, selected_course):
                    return True
        except Exception:
            pass

    return False


def _filter_attendance_students(program, course_id=None, verified='ALL', search=''):
    """Fetch admitted students for attendance sheet matching program and optional subject."""
    if not program:
        return [], None

    # Admitted students matching program either via Student or StudentAdmission
    adm_reg_nos = list(
        StudentAdmission.objects.filter(program_type=program)
        .exclude(status='Rejected')
        .values_list('reg_no', flat=True)
    )

    students_qs = Student.objects.filter(
        Q(program_type=program) | Q(registration_no__in=adm_reg_nos)
    )
    if verified == 'YES':
        students_qs = students_qs.filter(is_verified=True)
    elif verified == 'NO':
        students_qs = students_qs.filter(is_verified=False)
    if search:
        students_qs = students_qs.filter(
            Q(full_name__icontains=search) | Q(registration_no__icontains=search)
        )

    students = list(students_qs.order_by('registration_no', 'full_name'))
    students = _attach_admissions(students)

    # Exclude students whose admission was explicitly rejected
    students = [
        s for s in students
        if not getattr(s, 'latest_admission', None) or s.latest_admission.status != 'Rejected'
    ]

    selected_course = None
    if course_id and str(course_id).strip() and str(course_id).strip() != 'ALL':
        selected_course = ProgramCourse.objects.filter(pk=course_id).first()
        if selected_course and not selected_course.is_compulsory:
            students = [s for s in students if _student_takes_course(s, selected_course)]

    return students, selected_course


@admin_login_required
def attendance_sheets(request):
    """Admin module to preview and export Internal Assessment Attendance Sheets."""
    now = timezone.now()
    default_month = now.strftime('%B')
    default_year = str(now.year)

    program = request.GET.get('program', '').strip()
    course_id = request.GET.get('course_id', '').strip()
    month = request.GET.get('month', '').strip() or default_month
    year = request.GET.get('year', '').strip() or default_year
    assessment_title = request.GET.get('assessment_title', '').strip() or 'Internal Assessment'
    verified = request.GET.get('verified', 'ALL').strip() or 'ALL'
    search = request.GET.get('search', '').strip()

    program_types = get_program_names()
    available_courses = []
    students = []
    selected_course = None

    if program:
        available_courses = ProgramCourse.objects.filter(program_type=program).order_by('department', 'course_name', 'sort_order')
        if not available_courses.exists():
            base_prog = program.replace(' - First Semester', '').replace(' First Semester', '').strip()
            available_courses = ProgramCourse.objects.filter(
                Q(program_type=base_prog) | Q(program_type__istartswith=base_prog)
            ).order_by('department', 'course_name', 'sort_order')
        students, selected_course = _filter_attendance_students(program, course_id, verified, search)

    years_list = [str(now.year - 1), str(now.year), str(now.year + 1)]

    context = {
        'program': program,
        'course_id': course_id,
        'selected_course': selected_course,
        'month': month,
        'year': year,
        'assessment_title': assessment_title,
        'verified': verified,
        'search': search,
        'program_types': program_types,
        'available_courses': available_courses,
        'students': students,
        'total_students': len(students),
        'months_list': ATTENDANCE_MONTHS,
        'years_list': years_list,
    }
    return render(request, 'admin_panel/attendance_sheets.html', context)


@admin_login_required
def export_attendance_sheet_excel(request):
    """Export Internal Assessment Attendance Sheet strictly formatted for A4 Paper printing."""
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    now = timezone.now()
    program = request.GET.get('program', '').strip()
    if not program:
        messages.warning(request, 'Please select a program before exporting attendance sheet.')
        return redirect('attendance_sheets')

    course_id = request.GET.get('course_id', '').strip()
    month = request.GET.get('month', '').strip() or now.strftime('%B')
    year = request.GET.get('year', '').strip() or str(now.year)
    assessment_title = request.GET.get('assessment_title', '').strip() or 'Internal Assessment'
    verified = request.GET.get('verified', 'ALL').strip() or 'ALL'
    search = request.GET.get('search', '').strip()

    students, selected_course = _filter_attendance_students(program, course_id, verified, search)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Sheet"

    # Strict A4 Page Setup
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True
    ws.views.sheetView[0].showGridLines = True
    ws.print_title_rows = '1:7'

    thin_side = Side(style='thin', color='334155')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    thick_bottom_side = Side(style='medium', color='082B49')

    # Row 1: College Name
    ws.merge_cells('A1:E1')
    c1 = ws['A1']
    c1.value = "CHAITANYA SCIENCE AND ARTS COLLEGE, PAMGARH"
    c1.font = Font(name='Calibri', size=15, bold=True, color='082B49')
    c1.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 26

    # Row 2: Sub-affiliation
    ws.merge_cells('A2:E2')
    c2 = ws['A2']
    c2.value = "(An Autonomous Institution Approved by UGC | Affiliated to SNPV, Raigarh)"
    c2.font = Font(name='Calibri', size=9.5, italic=True, color='475569')
    c2.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 18

    # Row 3: Assessment Title & Month/Year
    ws.merge_cells('A3:E3')
    c3 = ws['A3']
    c3.value = f"{assessment_title.upper()} ATTENDANCE SHEET — {month.upper()} {year}"
    c3.font = Font(name='Calibri', size=12, bold=True, color='0F172A')
    c3.fill = PatternFill('solid', fgColor='F1F5F9')
    c3.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 23
    for col in range(1, 6):
        cell = ws.cell(row=3, column=col)
        cell.border = Border(top=thin_side, bottom=thin_side)

    # Row 4: Program & Session/Date
    ws.merge_cells('A4:C4')
    c4a = ws['A4']
    c4a.value = f"Class / Program: {program}"
    c4a.font = Font(name='Calibri', size=10.5, bold=True, color='1E293B')
    c4a.alignment = Alignment(horizontal='left', vertical='center')

    ws.merge_cells('D4:E4')
    c4b = ws['D4']
    c4b.value = "Session: 2026-27   |   Date: ______________"
    c4b.font = Font(name='Calibri', size=10, color='334155')
    c4b.alignment = Alignment(horizontal='right', vertical='center')
    ws.row_dimensions[4].height = 20

    # Row 5: Subject line (subject-wise vs general)
    ws.merge_cells('A5:E5')
    c5 = ws['A5']
    if selected_course:
        course_label = f"{selected_course.department} - {selected_course.course_name}" if selected_course.department else selected_course.course_name
        code_part = f" (Code: {selected_course.course_code})" if selected_course.course_code else ""
        c5.value = f"Subject / Course: {course_label}{code_part}"
        c5.font = Font(name='Calibri', size=11, bold=True, color='082B49')
        c5.fill = PatternFill('solid', fgColor='E0F2FE')
    else:
        c5.value = "Class Attendance Sheet (All Enrolled Subjects / General)"
        c5.font = Font(name='Calibri', size=10.5, italic=True, color='475569')
        c5.fill = PatternFill('solid', fgColor='F8FAFC')
    c5.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[5].height = 22
    for col in range(1, 6):
        ws.cell(row=5, column=col).border = Border(top=thin_side, bottom=thin_side)

    # Row 6: Spacer
    ws.row_dimensions[6].height = 6

    # Row 7: Table Headers
    headers = ['Sn', 'Student ID', 'NAME', 'Father Name', 'Signature']
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='082B49')
    header_align = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[7].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom_side)

    # Data Rows
    zebra_fill = PatternFill('solid', fgColor='F8FAFC')
    row_num = 8
    for idx, s in enumerate(students, start=1):
        ws.row_dimensions[row_num].height = 28  # Generous height for signing!
        fill_to_apply = zebra_fill if idx % 2 == 0 else PatternFill(fill_type=None)

        # Sn
        c_sn = ws.cell(row=row_num, column=1, value=idx)
        c_sn.alignment = Alignment(horizontal='center', vertical='center')
        c_sn.font = Font(name='Calibri', size=10)

        # Student ID
        c_id = ws.cell(row=row_num, column=2, value=s.registration_no)
        c_id.alignment = Alignment(horizontal='center', vertical='center')
        c_id.font = Font(name='Calibri', size=10, bold=True)

        # NAME
        c_name = ws.cell(row=row_num, column=3, value=s.full_name)
        c_name.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c_name.font = Font(name='Calibri', size=10, bold=True)

        # Father Name
        father = getattr(s, 'father_name', '')
        if not father and getattr(s, 'latest_admission', None):
            father = s.latest_admission.father_name or ''
        c_father = ws.cell(row=row_num, column=4, value=father or '—')
        c_father.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c_father.font = Font(name='Calibri', size=10)

        # Signature (Blank for physical handwriting)
        c_sig = ws.cell(row=row_num, column=5, value="")
        c_sig.alignment = Alignment(horizontal='center', vertical='center')

        for col_idx in range(1, 6):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border = cell_border
            if fill_to_apply.fill_type:
                cell.fill = fill_to_apply

        row_num += 1

    # Spacer row
    ws.row_dimensions[row_num].height = 10
    row_num += 1

    # Footer summary row
    ws.row_dimensions[row_num].height = 28
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    c_f1 = ws.cell(row=row_num, column=1, value=f"Total Students: {len(students)}       Present: _________       Absent: _________")
    c_f1.font = Font(name='Calibri', size=10, bold=True, color='082B49')
    c_f1.alignment = Alignment(horizontal='left', vertical='center')

    ws.merge_cells(start_row=row_num, start_column=4, end_row=row_num, end_column=5)
    c_f2 = ws.cell(row=row_num, column=4, value="Signature of Subject Teacher / Invigilator: __________________")
    c_f2.font = Font(name='Calibri', size=10, bold=True, color='082B49')
    c_f2.alignment = Alignment(horizontal='right', vertical='center')

    # Column widths formatted for A4 Portrait print
    ws.column_dimensions['A'].width = 6.5
    ws.column_dimensions['B'].width = 17.5
    ws.column_dimensions['C'].width = 28.5
    ws.column_dimensions['D'].width = 28.5
    ws.column_dimensions['E'].width = 26.5

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    prog_clean = re.sub(r'[^\w\-]+', '_', program).strip('_')
    sub_clean = f"_{re.sub(r'[^\w\-]+', '_', selected_course.course_name).strip('_')}" if selected_course else ""
    filename = f"Attendance_{prog_clean}{sub_clean}_{month}_{year}.xlsx"

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response